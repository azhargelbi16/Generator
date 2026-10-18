import pytest
from docx import Document
from generator.document import template


def _doc(path, lines):
    d = Document()
    for ln in lines:
        d.add_paragraph(ln)
    d.save(path)
    return path


def test_missing_reported(tmp_path):
    p = _doc(tmp_path / "t.docx", ["{{JUDUL}} {{NAMA}}"])
    res = template.validate_template(p)
    assert "{{DOSEN}}" in res.missing and "{{JUDUL}}" not in res.missing
    assert res.is_processable is False


def test_special_must_be_standalone(tmp_path):
    p = _doc(tmp_path / "t.docx", template.REQUIRED_PLACEHOLDERS + [])
    # ganti ISI_LAPORAN jadi tidak berdiri sendiri
    d = Document(p)
    for par in d.paragraphs:
        if par.text == "{{ISI_LAPORAN}}":
            par.text = "lihat {{ISI_LAPORAN}} di sini"
    d.save(p)
    res = template.validate_template(p)
    assert "{{ISI_LAPORAN}}" in res.not_standalone


def test_unknown_flagged_but_ok(tmp_path):
    p = _doc(tmp_path / "t.docx", template.REQUIRED_PLACEHOLDERS + ["{{FOO}}"])
    res = template.validate_template(p)
    assert "{{FOO}}" in res.unknown and res.is_processable is True


def test_all_present_is_processable(tmp_path):
    p = _doc(tmp_path / "t.docx", template.REQUIRED_PLACEHOLDERS)
    assert template.validate_template(p).is_processable is True


def test_bad_file_raises(tmp_path):
    p = tmp_path / "bad.docx"
    p.write_text("bukan docx")
    with pytest.raises(template.TemplateError):
        template.load_template(p)
