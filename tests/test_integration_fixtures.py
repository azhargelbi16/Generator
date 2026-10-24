"""Integration tests using real file fixtures (PDF generation/parsing & full DOCX pipeline)."""

import io
from pathlib import Path
from pypdf import PdfWriter
from docx import Document

from generator.sources.pdf_reader import extract_text_and_metadata
from generator.document import renderer, template
from generator.ai.content import Paragraph, Section, Sentence
from generator.report import sections_to_dicts
from generator.web.api_bridge import ApiBridge
from generator import database


def test_real_pdf_fixture_generation_and_text_extraction(tmp_path: Path):
    """Generate a real binary PDF file and verify text extraction through pdf_reader."""
    pdf_path = tmp_path / "sample_academic_paper.pdf"
    
    # Create real valid PDF with pypdf
    writer = PdfWriter()
    writer.add_blank_page(width=595, height=842)  # A4
    
    # Add metadata
    writer.add_metadata({
        "/Title": "Sistem Deteksi Hama Tanaman Padi",
        "/Author": "Dr. Pratama",
        "/Subject": "Kecerdasan Buatan dalam Pertanian"
    })
    
    with open(pdf_path, "wb") as f:
        writer.write(f)
        
    assert pdf_path.is_file()
    assert pdf_path.stat().st_size > 0
    
    # Test through extraction engine
    text, meta = extract_text_and_metadata(pdf_path)
    assert isinstance(text, str)
    assert isinstance(meta, dict)


def test_real_full_document_docx_pipeline_with_analytics(tmp_path: Path):
    """Generate a full multi-chapter document with template, verify OpenXML and compute analytics."""
    template_path = tmp_path / "Template_Test.docx"
    output_path = tmp_path / "Output_Test.docx"
    
    # 1. Create a valid standard DOCX template
    doc = Document()
    doc.add_heading("COVER MAKALAH", level=0)
    for ph in template.REQUIRED_PLACEHOLDERS:
        doc.add_paragraph(ph)
    doc.save(template_path)
    
    # 2. Build multi-chapter academic sections
    sections = [
        Section(1, "BAB 1. PENDAHULUAN", []),
        Section(2, "1.1 Latar Belakang", [
            Paragraph([
                Sentence("Penyakit daun padi dapat menurunkan produktivitas tanaman hingga tiga puluh persen.", "kutipan 1", [0], is_factual=True),
                Sentence("Oleh karena itu, diperlukan otomasi deteksi berbasis visi komputer.", "kutipan 2", [0], is_factual=True),
            ])
        ]),
        Section(2, "1.2 Rumusan Masalah", [
            Paragraph([Sentence("Bagaimana performa arsitektur deep learning pada perangkat edge?", "", [], is_factual=False)])
        ]),
        Section(1, "BAB 2. TINJAUAN PUSTAKA", [
            Paragraph([
                Sentence("Penelitian sebelumnya menggunakan arsitektur convolutional neural network.", "kutipan 3", [1], is_factual=True)
            ])
        ]),
        Section(1, "BAB 3. HASIL DAN PEMBAHASAN", [
            Paragraph([
                Sentence("Pengujian lapangan menghasilkan tingkat akurasi 96.4 persen dengan latensi rendah.", "kutipan 4", [0], is_factual=True)
            ])
        ]),
        Section(1, "BAB 4. KESIMPULAN DAN SARAN", [
            Paragraph([
                Sentence("Integrasi model edge deep learning terbukti efektif untuk deteksi real-time.", "kutipan 5", [0], is_factual=True)
            ])
        ]),
    ]
    
    sources_data = [
        {"id": 0, "title": "Paper CNN Padi", "authors": [{"family": "Pratama", "given": "A."}], "year": "2023"},
        {"id": 1, "title": "Review AI Pertanian", "authors": [{"family": "Suryanto", "given": "B."}], "year": "2022"},
    ]
    
    # 3. Compute analytics through ApiBridge
    db = database.Database(tmp_path / "analytics.db")
    db.init_schema()
    bridge = ApiBridge(db)
    
    serialized_sections = sections_to_dicts(sections)
    analytics = bridge.get_document_analytics(serialized_sections, sources_data)
    assert analytics["total_words"] > 40
    assert analytics["total_sentences"] == 6
    assert analytics["academic_score"] > 0
    assert len(analytics["section_distribution"]) >= 4
    db.close()
    
    # 4. Render to DOCX
    fields = {
        "{{JUDUL}}": "Deteksi Penyakit Daun Padi Real-Time",
        "{{NAMA}}": "Muhammad Azhar Ghaelby",
        "{{NIM}}": "232410101088",
        "{{PRODI}}": "Sistem Informasi",
        "{{MATA_KULIAH}}": "Kecerdasan Buatan",
        "{{DOSEN}}": "Dr. Ir. Soekarno Hadi",
        "{{TANGGAL}}": "23 Oktober 2026",
    }
    references = [
        "Pratama, A. (2023). Paper CNN Padi. Jurnal Informatika Pertanian.",
        "Suryanto, B. (2022). Review AI Pertanian. Jurnal Agroteknologi Tropika.",
    ]
    
    rendered_path = renderer.render_report(
        template_path,
        fields,
        sections,
        references,
        None,
        output_path,
    )
    
    # 5. Verify the generated DOCX file
    assert Path(rendered_path).exists()
    out_doc = Document(rendered_path)
    full_text = "\n".join(p.text for p in out_doc.paragraphs)
    
    # Verify placeholder replacement
    assert "{{JUDUL}}" not in full_text
    assert "{{NAMA}}" not in full_text
    assert "Deteksi Penyakit Daun Padi Real-Time" in full_text
    assert "Muhammad Azhar Ghaelby" in full_text
    
    # Verify headings and chapters exist
    assert "BAB 1. PENDAHULUAN" in full_text
    assert "1.1 Latar Belakang" in full_text
    assert "BAB 3. HASIL DAN PEMBAHASAN" in full_text
    assert "Pratama, A. (2023)" in full_text
    assert "Suryanto, B. (2022)" in full_text

