import base64

import pytest
from docx import Document
from docx.shared import Cm
from generator.document import renderer, template
from generator.ai.content import Sentence, Paragraph, Section


def test_save_gagal_tidak_merusak_docx_lama(tmp_path, monkeypatch):
    from pathlib import Path

    from docx.document import Document as DocClass

    out = tmp_path / "o.docx"
    out.write_bytes(b"lama")
    template_path = _tmpl(tmp_path)

    def rusak(self, path, *args, **kwargs):
        Path(path).write_bytes(b"rusak")
        raise OSError("disk penuh")

    monkeypatch.setattr(DocClass, "save", rusak)
    with pytest.raises(OSError):
        renderer.render_report(
            template_path,
            {ph: "x" for ph in template.REQUIRED_PLACEHOLDERS},
            [],
            [],
            None,
            out,
        )
    assert out.read_bytes() == b"lama"
    assert sorted(p.name for p in tmp_path.iterdir()) == ["o.docx", "t.docx"]


def _tmpl(tmp_path):
    p = tmp_path / "t.docx"
    d = Document()
    for ph in template.REQUIRED_PLACEHOLDERS:
        d.add_paragraph(ph)
    d.save(p)
    return p


def _png(path, w, h):
    png = base64.b64decode("iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mNkYAAAAAYAAjCB0C8AAAAASUVORK5CYII=")
    path.write_bytes(png)
    return path


def test_render_fills_fields_and_body(tmp_path):
    secs = [Section(1, "Pendahuluan", [Paragraph([Sentence("Isi bab.", "x", [0])])])]
    out = renderer.render_report(_tmpl(tmp_path),
        {"{{JUDUL}}": "Laporan X", "{{NAMA}}": "Budi", "{{NIM}}": "1",
         "{{MATA_KULIAH}}": "Agro", "{{DOSEN}}": "Dr. X", "{{TANGGAL}}": "2026"},
        secs, ["Santoso, B. (2020)."], None, tmp_path / "out.docx")
    doc = Document(out)
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "Laporan X" in text and "{{" not in text
    assert "Pendahuluan" in text and "Isi bab." in text and "Santoso, B. (2020)." in text
    assert any(p.style.name == "Heading 1" for p in doc.paragraphs)


def test_render_places_citations_from_labels(tmp_path):
    secs = [Section(1, "Bab", [Paragraph([Sentence("Pertanian organik meningkat pesat di Indonesia.", "q", [0])])])]
    out = renderer.render_report(_tmpl(tmp_path), {ph: "x" for ph in template.REQUIRED_PLACEHOLDERS},
                                 secs, ["Santoso, B. (2020)."], None, tmp_path / "o.docx",
                                 labels={0: "(Santoso, 2020)"})
    text = "\n".join(p.text for p in Document(out).paragraphs)
    assert "Pertanian organik meningkat pesat di Indonesia (Santoso, 2020)." in text


def test_toc_field_present(tmp_path):
    out = renderer.render_report(_tmpl(tmp_path), {ph: "x" for ph in template.REQUIRED_PLACEHOLDERS},
                                 [], [], None, tmp_path / "o.docx")
    assert "TOC" in Document(out).element.xml


def test_toc_mencakup_heading_sampai_level_sembilan(tmp_path):
    out = renderer.render_report(_tmpl(tmp_path), {ph: "x" for ph in template.REQUIRED_PLACEHOLDERS},
                                 [], [], None, tmp_path / "o.docx")
    assert 'TOC \\o "1-9"' in Document(out).element.xml


def test_render_menolak_template_dengan_penanda_khusus_tidak_standalone(tmp_path):
    path = tmp_path / "t.docx"
    doc = Document()
    for name in template.REQUIRED_PLACEHOLDERS:
        doc.add_paragraph("Bab: {{ISI_LAPORAN}}" if name == "{{ISI_LAPORAN}}" else name)
    doc.save(path)

    with pytest.raises(template.TemplateError, match="tidak siap"):
        renderer.render_report(
            path,
            {name: "x" for name in template.REQUIRED_PLACEHOLDERS},
            [],
            [],
            None,
            tmp_path / "o.docx",
        )


def test_render_menolak_template_tanpa_penanda_wajib(tmp_path):
    path = tmp_path / "t.docx"
    doc = Document()
    for name in template.REQUIRED_PLACEHOLDERS:
        if name != "{{DAFTAR_ISI}}":
            doc.add_paragraph(name)
    doc.save(path)

    with pytest.raises(template.TemplateError, match="DAFTAR_ISI"):
        renderer.render_report(
            path,
            {name: "x" for name in template.REQUIRED_PLACEHOLDERS},
            [],
            [],
            None,
            tmp_path / "o.docx",
        )


def test_special_placeholder_keys_in_fields_ignored(tmp_path):
    logo = _png(tmp_path / "l.png", 1, 1)
    fields = {ph: "x" for ph in template.REQUIRED_PLACEHOLDERS}
    out = renderer.render_report(_tmpl(tmp_path), fields, [], [], logo, tmp_path / "o.docx")
    doc = Document(out)
    assert len(doc.inline_shapes) == 1
    assert "TOC" in doc.element.xml


def test_logo_square_is_5cm(tmp_path):
    logo = _png(tmp_path / "l.png", 1, 1)
    out = renderer.render_report(_tmpl(tmp_path), {ph: "x" for ph in template.REQUIRED_PLACEHOLDERS},
                                 [], [], logo, tmp_path / "o.docx")
    shape = Document(out).inline_shapes[0]
    assert shape.width == Cm(5) and shape.height == Cm(5)


def test_logo_non_square_fits_within_5cm(tmp_path):
    from PIL import Image
    logo = tmp_path / "wide.png"; Image.new("RGB", (200, 100), "black").save(logo)
    out = renderer.render_report(_tmpl(tmp_path), {ph: "x" for ph in template.REQUIRED_PLACEHOLDERS},
                                 [], [], logo, tmp_path / "o.docx")
    shape = Document(out).inline_shapes[0]
    assert shape.width == Cm(5) and shape.height <= Cm(5)


def test_bad_logo_skipped(tmp_path):
    bad = tmp_path / "l.txt"; bad.write_text("nope")
    out = renderer.render_report(_tmpl(tmp_path), {ph: "x" for ph in template.REQUIRED_PLACEHOLDERS},
                                 [], [], bad, tmp_path / "o.docx")
    text = "\n".join(p.text for p in Document(out).paragraphs)
    assert out.exists() and "{{LOGO}}" not in text and "nope" not in text


def test_missing_heading_style_falls_back(tmp_path):
    p = tmp_path / "t.docx"
    d = Document()
    st = d.styles["Heading 1"].element
    st.getparent().remove(st)
    for ph in template.REQUIRED_PLACEHOLDERS: d.add_paragraph(ph)
    d.save(p)
    secs = [Section(1, "Bab Tanpa Style", [Paragraph([Sentence("Isi bab.", "x", [0])])])]
    out = renderer.render_report(p, {ph: "x" for ph in template.REQUIRED_PLACEHOLDERS},
                                  secs, [], None, tmp_path / "o.docx")
    text = "\n".join(par.text for par in Document(out).paragraphs)
    assert "Bab Tanpa Style" in text and "Isi bab." in text


def test_missing_optional_placeholder_is_blanked(tmp_path):
    p = tmp_path / "t.docx"
    d = Document()
    for ph in template.REQUIRED_PLACEHOLDERS: d.add_paragraph(ph)
    d.add_paragraph("Tahun {{TAHUN}} prodi {{PRODI}}")
    d.save(p)
    out = renderer.render_report(p, {ph: "x" for ph in template.REQUIRED_PLACEHOLDERS},
                                 [], [], None, tmp_path / "o.docx")
    text = "\n".join(par.text for par in Document(out).paragraphs)
    assert "{{" not in text and "Tahun" in text and "prodi" in text


def test_singkatan_organisasi_hanya_pada_kemunculan_pertama(tmp_path):
    panjang = "(Kementerian Pertanian [Kementan], 2020)"
    pendek = "(Kementan, 2020)"
    secs = [
        Section(
            1,
            "Bab",
            [
                Paragraph([Sentence("Satu.", "q", [0])]),
                Paragraph([Sentence("Dua.", "q", [0])]),
            ],
        )
    ]
    out = renderer.render_report(
        _tmpl(tmp_path),
        {ph: "x" for ph in template.REQUIRED_PLACEHOLDERS},
        secs,
        [],
        None,
        tmp_path / "o.docx",
        labels={0: panjang},
        label_forms={0: (panjang, pendek)},
    )
    texts = [p.text for p in Document(out).paragraphs]
    satu = next(t for t in texts if t.startswith("Satu "))
    dua = next(t for t in texts if t.startswith("Dua "))
    assert panjang in satu
    assert pendek in dua and panjang not in dua


def test_references_hanging_indent(tmp_path):
    out = renderer.render_report(_tmpl(tmp_path), {ph: "x" for ph in template.REQUIRED_PLACEHOLDERS},
                                 [], ["Santoso, B. (2020)."], None, tmp_path / "o.docx")
    doc = Document(out)
    ref = [p for p in doc.paragraphs if "Santoso, B. (2020)." in p.text][0]
    assert ref.paragraph_format.left_indent is not None
    assert ref.paragraph_format.first_line_indent is not None
    assert ref.paragraph_format.first_line_indent < 0


# --- Task 1.7: updateFields Word ---------------------------------------------


def _settings_xml(path) -> str:
    import zipfile

    with zipfile.ZipFile(path) as archive:
        return archive.read("word/settings.xml").decode("utf-8")


def test_update_fields_true_menambah_flag_update_fields(tmp_path):
    out = renderer.render_report(
        _tmpl(tmp_path), {ph: "x" for ph in template.REQUIRED_PLACEHOLDERS},
        [], [], None, tmp_path / "o.docx", update_fields=True,
    )
    assert 'w:updateFields w:val="true"' in _settings_xml(out)


def test_update_fields_default_tanpa_flag(tmp_path):
    out = renderer.render_report(
        _tmpl(tmp_path), {ph: "x" for ph in template.REQUIRED_PLACEHOLDERS},
        [], [], None, tmp_path / "o.docx",
    )
    assert "updateFields" not in _settings_xml(out)


def test_updatefields_ditaruh_sebelum_compat(tmp_path):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    path = _tmpl(tmp_path)
    doc = Document(path)
    doc.settings.element.append(OxmlElement("w:compat"))
    doc.settings.element.append(OxmlElement("w:rsids"))
    doc.save(path)

    out = renderer.render_report(
        path, {ph: "x" for ph in template.REQUIRED_PLACEHOLDERS}, [], [], None,
        tmp_path / "o.docx", update_fields=True,
    )
    tags = [child.tag for child in Document(out).settings.element]
    assert tags.index(qn("w:updateFields")) < tags.index(qn("w:compat"))


def test_logo_lama_tidak_tertinggal(tmp_path):
    from docx.oxml.ns import qn

    path = _tmpl(tmp_path)
    doc = Document(path)
    par = next(p for p in doc.paragraphs if p.text.strip() == "{{LOGO}}")
    par.add_run().add_picture(str(_png(tmp_path / "lama.png", 1, 1)), width=Cm(1))
    doc.save(path)

    out = renderer.render_report(
        path, {ph: "x" for ph in template.REQUIRED_PLACEHOLDERS}, [], [],
        _png(tmp_path / "logo.png", 2, 2), tmp_path / "o.docx",
    )
    drawings = list(Document(out).element.body.iter(qn("w:drawing")))
    assert len(drawings) == 1


def test_orientasi_exif_logo_dinormalkan(tmp_path):
    from PIL import Image

    path = tmp_path / "foto.jpg"
    image = Image.new("RGB", (40, 20), "red")
    exif = image.getexif()
    exif[274] = 6
    image.save(path, exif=exif)

    width, height = renderer._logo_size(path)
    assert height > width


def test_karakter_kontrol_dibersihkan(tmp_path):
    sections = [
        Section(1, "Bab\x0bSatu", [Paragraph([Sentence("Isi\x07 bab.", "x", [0])])])
    ]
    fields = {ph: "x" for ph in template.REQUIRED_PLACEHOLDERS}
    fields["{{JUDUL}}"] = "Laporan\x0cX"
    out = renderer.render_report(
        _tmpl(tmp_path), fields, sections, [], None, tmp_path / "o.docx"
    )
    text = "\n".join(p.text for p in Document(out).paragraphs)
    assert "BabSatu" in text
    assert "LaporanX" in text
    assert "\x0b" not in text and "\x07" not in text and "\x0c" not in text


def test_judul_berisi_penanda_tidak_menghapus_judul(tmp_path):
    sections = [Section(1, "Bab", [Paragraph([Sentence("Isi bab.", "x", [0])])])]
    fields = {ph: "x" for ph in template.REQUIRED_PLACEHOLDERS}
    fields["{{JUDUL}}"] = "{{ISI_LAPORAN}}"

    out = renderer.render_report(
        _tmpl(tmp_path), fields, sections, [], None, tmp_path / "o.docx"
    )
    texts = [p.text for p in Document(out).paragraphs]
    assert any("ISI_LAPORAN" in text for text in texts)
    assert sum("Isi bab." in text for text in texts) == 1


# --- Nomor halaman otomatis --------------------------------------------------


def test_nomor_halaman_ditambahkan_bila_footer_kosong(tmp_path):
    fields = {ph: "x" for ph in template.REQUIRED_PLACEHOLDERS}
    out = renderer.render_report(
        _tmpl(tmp_path), fields, [], [], None, tmp_path / "o.docx"
    )
    doc = Document(out)
    section = doc.sections[0]
    assert section.different_first_page_header_footer is True
    assert "PAGE" in section.footer._element.xml
    assert 'w:start="0"' in section._sectPr.xml


def test_footer_yang_sudah_berisi_teks_tidak_ditambahi_nomor(tmp_path):
    from docx import Document as DocxDocument

    path = tmp_path / "berfooter.docx"
    d = DocxDocument()
    for ph in template.REQUIRED_PLACEHOLDERS:
        d.add_paragraph(ph)
    d.sections[0].footer.paragraphs[0].text = "Prodi Agroteknologi"
    d.save(path)

    fields = {ph: "x" for ph in template.REQUIRED_PLACEHOLDERS}
    out = renderer.render_report(path, fields, [], [], None, tmp_path / "o.docx")
    doc = Document(out)
    footer_text = "\n".join(p.text for p in doc.sections[0].footer.paragraphs)
    assert "Prodi Agroteknologi" in footer_text
    assert "PAGE" not in doc.sections[0].footer._element.xml


# --- Satu butir satu paragraf & daftar pustaka italic ------------------------


def _fields():
    return {ph: "x" for ph in template.REQUIRED_PLACEHOLDERS}


def test_butir_bernomor_dipecah_menjadi_paragraf_terpisah(tmp_path):
    sections = [Section(1, "Bab", [Paragraph([
        Sentence("1. Bagaimana pengaruh media sosial?", "x", [0]),
        Sentence("2. Bagaimana peran waktu penggunaan?", "x", [0]),
        Sentence("3. Bagaimana dampaknya pada prestasi?", "x", [0]),
    ])])]
    out = renderer.render_report(
        _tmpl(tmp_path), _fields(), sections, [], None, tmp_path / "o.docx"
    )
    texts = [p.text for p in Document(out).paragraphs]
    idx = [next(i for i, t in enumerate(texts) if t.strip().startswith(f"{n}. ")) for n in (1, 2, 3)]
    assert idx[1] == idx[0] + 1
    assert idx[2] == idx[1] + 1


def test_paragraf_naratif_tetap_satu_paragraf(tmp_path):
    sections = [Section(1, "Bab", [Paragraph([
        Sentence("Media sosial telah menjadi bagian hidup mahasiswa.", "x", [0]),
        Sentence("Penggunaannya berdampak pada prestasi akademik.", "x", [0]),
        Sentence("Dampak itu bergantung pada motif penggunaan.", "x", [0]),
    ])])]
    out = renderer.render_report(
        _tmpl(tmp_path), _fields(), sections, [], None, tmp_path / "o.docx"
    )
    texts = [p.text for p in Document(out).paragraphs]
    gabungan = [t for t in texts if "Media sosial telah menjadi" in t]
    assert len(gabungan) == 1
    assert "Dampak itu bergantung" in gabungan[0]


def test_referensi_memakai_run_italic_untuk_nama_jurnal(tmp_path):
    sections = [Section(1, "Bab", [Paragraph([Sentence("Isi bab.", "x", [0])])])]
    refs = [[
        ("Santoso, B. (2024). Judul. ", False),
        ("Healthcare, 9", True),
        ("(10), 1272.", False),
    ]]
    out = renderer.render_report(
        _tmpl(tmp_path), _fields(), sections, refs, None, tmp_path / "o.docx"
    )
    doc = Document(out)
    italic_runs = [
        run.text for p in doc.paragraphs for run in p.runs if run.italic
    ]
    assert any("Healthcare, 9" in text for text in italic_runs)


def test_referensi_string_biasa_tetap_dirender(tmp_path):
    sections = [Section(1, "Bab", [Paragraph([Sentence("Isi bab.", "x", [0])])])]
    out = renderer.render_report(
        _tmpl(tmp_path), _fields(), sections,
        ["Santoso, B. (2020). Judul. Jurnal Tanah."], None, tmp_path / "o.docx",
    )
    texts = [p.text for p in Document(out).paragraphs]
    assert any("Jurnal Tanah." in text for text in texts)



def test_judul_cover_otomatis_title_case(tmp_path):
    sections = [Section(1, "Bab Satu", [Paragraph([Sentence("Isi bab.", "x", [0])])])]
    fields = {ph: "x" for ph in template.REQUIRED_PLACEHOLDERS}
    fields["{{JUDUL}}"] = "dampak inovasi teknologi pada pertanian"
    fields["{{DOSEN}}"] = "Dr. Ir. Soetomo, M.P."
    out = renderer.render_report(
        _tmpl(tmp_path), fields, sections, [], None, tmp_path / "title_case.docx"
    )
    text = "\n".join(p.text for p in Document(out).paragraphs)
    assert "Dampak Inovasi Teknologi pada Pertanian" in text


def test_dosen_pengampu_dihapus_jika_kosong(tmp_path):
    doc = Document()
    doc.add_paragraph("{{JUDUL}}")
    doc.add_paragraph("{{NAMA}}")
    doc.add_paragraph("{{NIM}}")
    doc.add_paragraph("{{MATA_KULIAH}}")
    doc.add_paragraph("Dosen Pengampu:")
    doc.add_paragraph("{{DOSEN}}")
    doc.add_paragraph("{{TANGGAL}}")
    doc.add_paragraph("{{DAFTAR_ISI}}")
    doc.add_paragraph("{{LOGO}}")
    doc.add_paragraph("{{ISI_LAPORAN}}")
    doc.add_paragraph("{{DAFTAR_PUSTAKA}}")
    tmpl_path = tmp_path / "tmpl_dosen.docx"
    doc.save(tmpl_path)

    sections = [Section(1, "Bab Satu", [Paragraph([Sentence("Isi.", "x", [0])])])]
    fields = {ph: "x" for ph in template.REQUIRED_PLACEHOLDERS}
    fields["{{DOSEN}}"] = ""
    fields["DOSEN"] = ""
    out = renderer.render_report(
        tmpl_path, fields, sections, [], None, tmp_path / "no_dosen.docx"
    )
    paragraphs = [p.text.strip() for p in Document(out).paragraphs]
    assert not any("Dosen Pengampu:" in p for p in paragraphs)
    assert not any("Dosen Pembimbing:" in p for p in paragraphs)

    # Bila diisi, label dan nama dosen tetap muncul
    fields["{{DOSEN}}"] = "Prof. Dr. Ir. Budi"
    fields["DOSEN"] = "Prof. Dr. Ir. Budi"
    out2 = renderer.render_report(
        tmpl_path, fields, sections, [], None, tmp_path / "with_dosen.docx"
    )
    paragraphs2 = [p.text.strip() for p in Document(out2).paragraphs]
    assert any("Dosen Pengampu:" in p for p in paragraphs2)
    assert any("Prof. Dr. Ir. Budi" in p for p in paragraphs2)
