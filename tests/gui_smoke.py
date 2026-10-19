"""Smoke test GUI Task 13: jendela terbuka, tab ada, tambah dosen tersimpan lintas aplikasi."""

import sys
import tempfile
import tkinter as ttk_tk  # noqa: F401  (pastikan Tk tersedia sebelum import app)
from tkinter import ttk

from generator import paths

paths.set_data_dir(tempfile.mkdtemp(prefix="generator-smoke-"))
paths.ensure_dirs()

from generator.gui import app  # noqa: E402
from generator.database import Database  # noqa: E402


def find_notebook(widget):
    for child in widget.winfo_children():
        if isinstance(child, ttk.Notebook):
            return child
        found = find_notebook(child)
        if found is not None:
            return found
    return None


# 1) Jendela pertama: cek tab, tambah satu dosen lewat panel GUI
root = app.build_app()
assert root is not None, "build_app mengembalikan None"
root.update()
nb = find_notebook(root)
assert nb is not None, "Notebook tidak ditemukan"
tab_ids = nb.tabs()
names = [nb.tab(t, "text") for t in tab_ids]
assert names == ["Data", "Template", "Laporan", "Review Hasil", "Pengaturan"], names

nb.select(tab_ids[0])
root.update()
window_bottom = root.winfo_rooty() + root.winfo_height()
data_tab = nb.nametowidget(tab_ids[0])
inner = find_notebook(data_tab)
assert inner is not None, "Notebook dalam DataTab tidak ditemukan"
inner_names = [inner.tab(t, "text") for t in inner.tabs()]
assert inner_names == ["Profil Mahasiswa", "Dosen", "Mata Kuliah", "Logo"], inner_names


def assert_buttons_visible(panel, label):
    for btn in panel.winfo_children():
        pass
    stack = list(panel.winfo_children())
    while stack:
        widget = stack.pop()
        stack.extend(widget.winfo_children())
        if isinstance(widget, ttk.Button) and widget.winfo_ismapped():
            bottom = widget.winfo_rooty() + widget.winfo_height()
            assert bottom <= window_bottom, (
                f"Tombol '{widget.cget('text')}' di {label} terpotong "
                f"(bawah {bottom} > jendela {window_bottom})"
            )


for sub_index, sub_name in [(1, "Dosen"), (2, "Mata Kuliah"), (3, "Logo")]:
    inner.select(sub_index)
    root.update()
    panel = inner.nametowidget(inner.tabs()[sub_index])
    assert_buttons_visible(panel, sub_name)

dosen_panel = inner.nametowidget(inner.tabs()[1])
dosen_panel._vars["nama"].set("Siti Aminah")
dosen_panel._vars["gelar"].set("S.P.")
dosen_panel._vars["identitas"].set("19900101")
dosen_panel._add()
root.update()
assert dosen_panel._listbox.size() == 1, "Dosen tidak muncul di listbox"

matkul_panel = inner.nametowidget(inner.tabs()[2])
matkul_panel._vars["nama"].set("Dasar Pertanian")
matkul_panel._vars["kode"].set("AGR101")
matkul_panel._add()
root.update()
assert matkul_panel._listbox.size() == 1, "Mata kuliah tidak muncul di listbox"

root.destroy()  # tutup aplikasi pertama

# 2) Jendela kedua: data harus masih ada setelah aplikasi dibuka ulang
root2 = app.build_app()
assert root2 is not None
root2.update()
db = Database(paths.DB_PATH)
db.init_schema()
dosen = db.list_dosen()
matkul = db.list_mata_kuliah()
db.close()
assert len(dosen) == 1 and dosen[0].nama == "Siti Aminah", dosen
assert len(matkul) == 1 and matkul[0].kode == "AGR101", matkul
# 3) Tab Pengaturan: bukan placeholder, Simpan menulis config.json
from generator.gui.tab_settings import SettingsTab  # noqa: E402
from generator.config import load_config  # noqa: E402

root2.update()
settings_tab = nb2 = None
main_nb2 = find_notebook(root2)
tab_ids2 = main_nb2.tabs()
main_nb2.select(tab_ids2[4])
root2.update()
settings_tab = main_nb2.nametowidget(tab_ids2[4])
assert isinstance(settings_tab, SettingsTab), type(settings_tab)
assert_buttons_visible(settings_tab, "Pengaturan")

settings_tab._base_url.set("http://localhost:9999/v1")
settings_tab._api_key.set("kunci-uji")
settings_tab._model.set("model-uji")
settings_tab._temperature.set("0.7")
settings_tab._save()
root2.update()
assert settings_tab._status.cget("text") == "Pengaturan tersimpan.", settings_tab._status.cget("text")
saved = load_config(paths.CONFIG_PATH)
assert saved.ai_base_url == "http://localhost:9999/v1", saved
assert saved.ai_api_key == "kunci-uji", saved
assert saved.ai_model == "model-uji", saved
assert abs(saved.ai_temperature - 0.7) < 1e-9, saved

# Validasi suhu waktu nyata: petunjuk berubah merah saat di luar rentang, pulih saat valid
settings_tab._temperature.set("5")
# Event keyboard butuh fokus OS (tidak tersedia di jendela latar); FocusOut
# tidak bergantung fokus dan memicu penangan yang sama, persis seperti user
# pindah bidang setelah mengetik.
settings_tab._entry_temperatur.event_generate("<FocusOut>")
root2.update()
assert "0 dan 2" in settings_tab._petunjuk_temperatur.cget("text"), settings_tab._petunjuk_temperatur.cget("text")
assert str(settings_tab._petunjuk_temperatur.cget("foreground")) == "#b00020", settings_tab._petunjuk_temperatur.cget("foreground")
settings_tab._temperature.set("0.7")
settings_tab._entry_temperatur.event_generate("<FocusIn>")
settings_tab._entry_temperatur.event_generate("<FocusOut>")
root2.update()
assert str(settings_tab._petunjuk_temperatur.cget("foreground")) == "#666666", settings_tab._petunjuk_temperatur.cget("foreground")

# Rentang di luar 0-2 tidak boleh menyimpan
settings_tab._temperature.set("2.5")
settings_tab._save()
root2.update()
assert "0 dan 2" in settings_tab._status.cget("text"), settings_tab._status.cget("text")
settings_tab._temperature.set("0.7")

# Tes Koneksi terkunci saat menguji: tombol mati berlabel Menguji..., pulih setelahnya
from generator.ai import client as klien_modul  # noqa: E402
import time as _time  # noqa: E402

_asli_check = klien_modul.check_connection

def check_lambat(_ai):
    _time.sleep(0.6)
    return True, "ok"

klien_modul.check_connection = check_lambat
settings_tab._test()
root2.update()
assert str(settings_tab._tombol_tes.cget("state")) == "disabled", settings_tab._tombol_tes.cget("state")
assert settings_tab._tombol_tes.cget("text") == "Menguji...", settings_tab._tombol_tes.cget("text")
mulai = _time.time()
while _time.time() - mulai < 5:
    root2.update()
    if str(settings_tab._tombol_tes.cget("state")) == "normal":
        break
    _time.sleep(0.05)
assert str(settings_tab._tombol_tes.cget("state")) == "normal", settings_tab._tombol_tes.cget("state")
assert settings_tab._tombol_tes.cget("text") == "Tes Koneksi", settings_tab._tombol_tes.cget("text")
assert "berhasil" in settings_tab._status.cget("text"), settings_tab._status.cget("text")
klien_modul.check_connection = _asli_check


# Temperatur salah format tidak boleh menyimpan
settings_tab._temperature.set("bukan-angka")
settings_tab._save()
rejected = load_config(paths.CONFIG_PATH)
assert rejected.ai_temperature == 0.7, rejected

# 4) Tab Template: validasi saat unggah, status dihitung ulang saat daftar tampil
import generator.gui.tab_template as tab_template_module  # noqa: E402
from pathlib import Path  # noqa: E402
from docx import Document  # noqa: E402
from generator.document.template import (  # noqa: E402
    OPTIONAL_PLACEHOLDERS,
    REQUIRED_PLACEHOLDERS,
    SPECIAL_PLACEHOLDERS,
)

dialogs: dict[str, list] = {"warning": [], "error": [], "askyesno": []}
tab_template_module.messagebox.showwarning = lambda *a, **k: dialogs["warning"].append(a)
tab_template_module.messagebox.showerror = lambda *a, **k: dialogs["error"].append(a)
tab_template_module.messagebox.askyesno = lambda *a, **k: dialogs["askyesno"].append(a) or True

main_nb2.select(tab_ids2[1])
root2.update()
template_tab = main_nb2.nametowidget(tab_ids2[1])
assert_buttons_visible(template_tab, "Template")

tmp = Path(tempfile.mkdtemp(prefix="template-uji-"))

complete = tmp / "lengkap.docx"
doc = Document()
doc.add_paragraph("JUDUL LAPORAN {{JUDUL}}")
for name in SPECIAL_PLACEHOLDERS:
    doc.add_paragraph(f"teks sebelum {name} teks sesudah")
for name in REQUIRED_PLACEHOLDERS + OPTIONAL_PLACEHOLDERS:
    doc.add_paragraph(name)
doc.save(complete)

assert template_tab._import_template(complete, nama="Template Lengkap", deskripsi="uji"), (
    "Template lengkap gagal disimpan"
)
assert not dialogs["warning"], dialogs["warning"]
assert "lengkap" in template_tab._listbox.get(0), template_tab._listbox.get(0)

# Status dihitung ulang saat tampil: rusakkan SALINAN yang tersimpan lalu refresh.
incomplete = Document()
incomplete.add_paragraph("Hanya {{JUDUL}}")
incomplete_doc = tmp / "tidak_lengkap.docx"
incomplete.save(incomplete_doc)
stored = Path(template_tab._items[0].path)
incomplete.save(stored)
template_tab.refresh()
assert "tidak lengkap" in template_tab._listbox.get(0), template_tab._listbox.get(0)

# Unggah template tidak lengkap: peringatan tapi tetap tersimpan.
assert template_tab._import_template(incomplete_doc, nama="Template Parcial"), (
    "Template tidak lengkap harus tetap disimpan"
)
assert dialogs["warning"], "Peringatan tidak lengkap tidak muncul"
assert "kurang" in str(dialogs["warning"][0][1]), dialogs["warning"]

# Berkas rusak: dialog error, tidak tersimpan.
broken = tmp / "rusak.docx"
broken.write_text("ini bukan docx", encoding="utf-8")
assert not template_tab._import_template(broken, nama="Rusak")
assert dialogs["error"], "Dialog error untuk berkas rusak tidak muncul"

# Hapus terpilih.
db2 = Database(paths.DB_PATH)
jumlah_sebelum = len(db2.list_template())
db2.close()
template_tab._listbox.selection_set(0)
template_tab._delete()
db2 = Database(paths.DB_PATH)
jumlah_sesudah = len(db2.list_template())
db2.close()
assert jumlah_sesudah == jumlah_sebelum - 1, (jumlah_sebelum, jumlah_sesudah)
assert dialogs["askyesno"], "Konfirmasi hapus tidak muncul"

# 5) Tab Laporan: sumber manual, koreksi metadata, validasi template, outline, draf
import json  # noqa: E402
import shutil  # noqa: E402
import time  # noqa: E402

import generator.gui.tab_report as tab_report_module  # noqa: E402
from generator.ai.citation import Author, Source  # noqa: E402
from generator.database import Template  # noqa: E402
from generator.gui.tab_report import normalize_metadata  # noqa: E402


def pump(root, predicate, timeout_s=8.0):
    deadline = time.time() + timeout_s
    while time.time() < deadline:
        root.update()
        if predicate():
            return True
        time.sleep(0.02)
    return False


# Sumber lengkap masuk DB sebelum jendela dibangun, supaya Combobox ikut memuatnya.
final_src = tmp / "final.docx"
doc_final = Document()
doc_final.add_paragraph("JUDUL {{JUDUL}}")
for name in SPECIAL_PLACEHOLDERS:
    doc_final.add_paragraph(f"teks sebelum {name} teks sesudah")
for name in REQUIRED_PLACEHOLDERS + OPTIONAL_PLACEHOLDERS:
    doc_final.add_paragraph(name)
doc_final.save(final_src)
final_copy = paths.TEMPLATES_DIR / "final.docx"
shutil.copy2(final_src, final_copy)
db3 = Database(paths.DB_PATH)
db3.add_template(Template(0, "Template Final", "uji", str(final_copy)))
db3.close()

root3 = app.build_app()
assert root3 is not None
root3.update()
nb3 = find_notebook(root3)
tab_ids3 = nb3.tabs()
nb3.select(tab_ids3[2])
root3.update()
report_tab = nb3.nametowidget(tab_ids3[2]).controller
assert type(report_tab).__name__ == "ReportController", type(report_tab).__name__
assert_buttons_visible(nb3.nametowidget(tab_ids3[2]), "Laporan")

assert normalize_metadata({"date": "2021-05"})["year"] == "2021"
assert normalize_metadata({"date": "D:20231019120000+07'00'"})["year"] == "2023"
assert "year" not in normalize_metadata({"title": "X"})

# Sumber manual: extract_source dipalsukan agar tidak butuh layanan AI.
canned = Source(
    authors=[Author("Santoso", "Budi"), Author("Kementerian", is_org=True, abbrev="KP")],
    year="2020",
    title="Laporan Pertanian",
    url="http://contoh/id",
)
tab_report_module.bibliography.extract_source = lambda text, client: canned
report_tab._add_manual_text("Teks sumber manual untuk pengujian.")
assert pump(root3, lambda: len(report_tab._sources) == 1), "Sumber manual tidak masuk"
record = report_tab._sources[0]
assert record["kind"] == "manual" and record["title"] == "Laporan Pertanian", record
assert record["text"] == "Teks sumber manual untuk pengujian."
draft_file = paths.DRAFTS_DIR / "aktif.json"
assert draft_file.is_file(), "Draf tidak ditulis"
draft_raw = json.loads(draft_file.read_text(encoding="utf-8"))
assert draft_raw["sources"][0]["text"].startswith("Teks sumber manual"), draft_raw.keys()

# Regresi: kedua dialog sumber tidak crash saat dipanggil seperti klik user
# (pernah rusak karena parent dialog bukan widget Tk).
asli_askstring = tab_report_module.simpledialog.askstring
tab_report_module.simpledialog.askstring = lambda *a, **k: ""
report_tab._add_link()  # URL kosong: berhenti setelah dialog, tanpa jaringan
tab_report_module.simpledialog.askstring = lambda *a, **k: "https://contoh.go.id/artikel"
tujuan_link = {}
report_tab._add_link_core = lambda url: tujuan_link.setdefault("url", url)
report_tab._add_link()
assert tujuan_link.get("url") == "https://contoh.go.id/artikel", tujuan_link
tab_report_module.simpledialog.askstring = asli_askstring

report_tab._paste_manual()
root3.update()
tops = [w for w in report_tab._widget.winfo_children() if isinstance(w, ttk_tk.Toplevel)]
assert tops, "Dialog Tempel Teks Manual tidak terbuka"
for top in tops:
    top.destroy()
root3.update()

# Koreksi metadata diterapkan ke record dan draf.
report_tab._meta_vars["year"].set("2019")
report_tab._meta_vars["title"].set("Judul Koreksi")
report_tab._author_vars[0]["family"].set("SantosoKoreksi")
report_tab._apply_corrections()
draft_raw = json.loads(draft_file.read_text(encoding="utf-8"))
assert draft_raw["sources"][0]["year"] == "2019", draft_raw["sources"][0]
assert draft_raw["sources"][0]["authors"][0]["family"] == "SantosoKoreksi"

# Validasi template saat dipilih: tidak lengkap ditolak, lengkap diterima.
warning_before = len(dialogs["warning"])
report_tab._template_var.set("Template Parcial")
report_tab._on_template_selected()
assert len(dialogs["warning"]) == warning_before + 1, "Template tidak lengkap tidak diperingatkan"
assert report_tab._template is None and report_tab._template_var.get() == ""

report_tab._template_var.set("Template Final")
report_tab._on_template_selected()
assert report_tab._template is not None, "Template lengkap harus diterima"
assert len(dialogs["warning"]) == warning_before + 1, "Template lengkap tak boleh memunculkan peringatan"

# Usulkan Outline (AI dipalsukan) mengisi kotak outline dan menyembunyikan tombol.
report_tab._judul_var.set("Pertanian Organik")
report_tab._outline_text.delete("1.0", "end")
tab_report_module.content.suggest_outline = (
    lambda topic, sources, texts, client: ["Pendahuluan", "Tinjauan Pustaka"]
)
report_tab._suggest_outline()
assert pump(root3, lambda: report_tab._outline_content() == ["Pendahuluan", "Tinjauan Pustaka"]), (
    "Outline tidak terisi"
)
assert not report_tab._suggest_button.winfo_ismapped(), "Tombol usulan harus hilang saat outline terisi"

# Generate isi laporan (AI dipalsukan): status per kalimat + gerbang ekspor + draf.
from generator.ai.content import Paragraph, Section, Sentence  # noqa: E402
from generator import report as report_module  # noqa: E402

report_tab._add_manual_text(
    "Teks sumber manual kedua. Pertanian organik di indonesia sangat penting "
    "bagi petani di musim hujan tahun ini."
)
assert pump(root3, lambda: len(report_tab._sources) == 2), "Sumber kedua tidak masuk"

# Tombol hapus sumber harus tombol tk berwarna tegas; gaya ttk kustom tidak
# digambar sv-ttk (teks putih di latar putih, tidak terlihat sama sekali).
tombol_hapus = report_tab._tombol_hapus
assert isinstance(tombol_hapus, ttk_tk.Button), "Tombol hapus harus tk.Button berwarna tegas"
assert tombol_hapus.cget("background") == "#c0392b", tombol_hapus.cget("background")
assert tombol_hapus.cget("foreground") == "#ffffff", tombol_hapus.cget("foreground")

canned_section = Section(
    1,
    "Pendahuluan",
    [
        Paragraph(
            [
                Sentence(
                    "Kalimat terverifikasi.",
                    "pertanian organik di indonesia sangat penting bagi petani di musim hujan",
                    [1],
                ),
                Sentence("Bab ini membahas latar belakang.", is_factual=False),
                Sentence(
                    "Kalimat belum terverifikasi.",
                    "kata kata kata kata kata kata kata kata",
                    [1],
                ),
            ]
        )
    ],
)
tab_report_module.content.generate_sections = lambda *args, **kwargs: [canned_section]
report_tab._generate()
assert pump(root3, lambda: len(report_tab._statuses) == 3), "Status kalimat tidak terisi"
assert pump(root3, lambda: len(report_tab._sentence_ranges) == 3), "Kalimat tidak terbangun di dokumen"
status_dok = {report_tab._uid_status(u) for u in report_tab._sentence_ranges}
assert status_dok == {"terverifikasi", "tidak_terverifikasi", "bukan_klaim_faktual"}, status_dok
assert "Sudah selesai diverifikasi" in report_tab._chip_status.cget("text"), report_tab._chip_status.cget("text")
assert "dari 3 kalimat" in report_tab._chip_status.cget("text"), report_tab._chip_status.cget("text")

# Tampilan dokumen bersih ala Word: TNR 12, tanpa warna status, sitasi inline
isi_dok = report_tab._doc.get("1.0", "end-1c")
assert "\n\n" in isi_dok, "Paragraf tidak dipisah baris kosong"
assert report_tab._doc.cget("font") == "{Times New Roman} 12", report_tab._doc.cget("font")
assert report_tab._doc.cget("width") == 92, report_tab._doc.cget("width")
semua_tag = set()
for uid in report_tab._sentence_ranges:
    semua_tag.update(report_tab._doc.tag_names(report_tab._sentence_ranges[uid][0]))
assert not ({"terverifikasi", "tidak_terverifikasi", "bukan_klaim_faktual"} & semua_tag), semua_tag
assert "Santoso" in isi_dok and "2020" in isi_dok, "Sitasi inline tidak tercetak di dokumen"
assert not hasattr(report_tab, "_tombol_tema"), "Fitur mode gelap harus dihapus"

ok, message = report_module.can_export(report_tab._statuses)
assert ok is False and message, "Gerbang ekspor harus memblokir kalimat belum selesai"
pending_uids = [
    report_tab._statuses[item].sentence.uid
    for item in report_module.pending(report_tab._statuses, frozenset())
]
assert len(pending_uids) == 2, pending_uids

# Pulihkan sitasi pada kalimat BELUM lewat pemilihan di dokumen.
belum_uid = next(
    u for u in report_tab._sentence_ranges if report_tab._uid_status(u) == "tidak_terverifikasi"
)
report_tab._sentence_by_uid(belum_uid).citation = False
report_tab._select_sentence(belum_uid)
root3.update()
assert report_tab._selected_uid == belum_uid, "Klik kalimat tidak memilihnya"
report_tab._restore_citation_action()
assert report_tab._sentence_by_uid(belum_uid).citation is True, "Sitasi tidak dipulihkan"

# Klik kalimat: baris aksi melayang muncul di dokumen dengan status + tombol
report_tab._select_sentence(belum_uid)
root3.update()
assert report_tab._aksi_frame is not None and report_tab._aksi_frame.winfo_exists(), (
    "Baris aksi tidak muncul"
)
tombol = [
    w.cget("text") for w in report_tab._aksi_frame.winfo_children() if w.winfo_class() == "TButton"
]
assert "✒ Parafrase" in tombol and "✔ Tandai selesai" in tombol, tombol

# Klik area kosong: baris aksi hilang, sorotan lepas
report_tab._deselect()
root3.update()
assert report_tab._aksi_frame is None, "Baris aksi harus hilang saat lepas"
assert report_tab._selected_uid is None, "Pilihan harus lepas saat klik kosong"

# Sunting langsung di dokumen: ketik di rentang kalimat, tersimpan saat lepas.
report_tab._select_sentence(belum_uid)
awal, akhir = report_tab._sentence_ranges[belum_uid]
report_tab._doc.delete(awal, akhir)
report_tab._doc.insert(awal, "Kalimat diedit langsung di dokumen.")
report_tab._commit_inline_edit()
edited = report_tab._sentence_by_uid(belum_uid)
assert edited.origin == "manual", "Suntingan langsung harus menandai kalimat manual"
assert edited.uid == belum_uid, "Suntingan harus mempertahankan uid"
assert edited.text == "Kalimat diedit langsung di dokumen.", edited.text
report_tab._select_sentence(belum_uid)
report_tab._mark_done()

# Tandai semua bukan klaim faktual selesai lewat tombol.
nonfactual_uid = next(
    u for u in report_tab._sentence_ranges if report_tab._uid_status(u) == "bukan_klaim_faktual"
)
report_tab._select_sentence(nonfactual_uid)
report_tab._mark_all_nonfactual_done()
assert report_module.can_export(report_tab._statuses, frozenset(report_tab._done)) == (True, ""), (
    "Tanda selesai harus membuka gerbang ekspor"
)

# Parafrase kalimat terpilih (AI dipalsukan): kalimat baru uid segar, status dihitung ulang.
def fake_paraphrase(paragraph, sources, texts, client):
    return Paragraph(
        [
            Sentence(
                "Parafrase uji.",
                "pertanian organik di indonesia sangat penting bagi petani di musim hujan",
                [1],
            )
        ]
    )

tab_report_module.content.paraphrase = fake_paraphrase
report_tab._select_sentence(belum_uid)
root3.update()
report_tab._paraphrase()
assert pump(root3, lambda: any(s.text == "Parafrase uji." for s in report_tab._iter_sentences())), (
    "Parafrase tidak mengganti kalimat"
)
para_uid = next(s.uid for s in report_tab._iter_sentences() if s.text == "Parafrase uji.")
assert para_uid != belum_uid, "Parafrase harus menghasilkan uid baru"
assert report_tab._uid_status(para_uid) == "terverifikasi", (
    "Parafrase dengan kutipan sah harus terverifikasi"
)
assert report_module.can_export(report_tab._statuses, frozenset(report_tab._done)) == (True, "")

# Ketikan di luar kalimat aktif ditolak; di dalam diterima.
report_tab._select_sentence(para_uid)
awal, akhir = report_tab._sentence_ranges[para_uid]
kelas = ttk_tk.Event()
kelas.keysym = "a"
kelas.state = 0
report_tab._doc.mark_set("insert", f"{awal} linestart")
assert report_tab._on_doc_key(kelas) == "break", "Ketikan di luar kalimat aktif harus diblok"
report_tab._doc.mark_set("insert", awal)
assert report_tab._on_doc_key(kelas) is None, "Ketikan di dalam kalimat aktif harus boleh"

# Enter membelah kalimat; BackSpace/Delete di ujung rentang mengenai karimat lain.
# Ketiganya harus diblok meski kursor masih di dalam atau di tepi rentang.
kelas.keysym = "Return"
report_tab._doc.mark_set("insert", awal)
assert report_tab._on_doc_key(kelas) == "break", "Enter harus diblok di dalam kalimat aktif"
kelas.keysym = "BackSpace"
report_tab._doc.mark_set("insert", awal)
assert report_tab._on_doc_key(kelas) == "break", "BackSpace di tepi awal harus diblok"
kelas.keysym = "Delete"
report_tab._doc.mark_set("insert", akhir)
assert report_tab._on_doc_key(kelas) == "break", "Delete di tepi akhir harus diblok"

# Seleksi yang melintasi batas kalimat ditolak walau kursor di dalam rentang.
report_tab._doc.tag_add("sel", f"{awal}+2c", f"{akhir}+2c")
kelas.keysym = "a"
report_tab._doc.mark_set("insert", f"{awal}+3c")
assert report_tab._on_doc_key(kelas) == "break", "Seleksi melintasi batas kalimat harus diblok"
report_tab._doc.tag_remove("sel", "1.0", "end")
report_tab._deselect()

draft_raw = json.loads(draft_file.read_text(encoding="utf-8"))
assert len(draft_raw["sections"]) == 1, draft_raw["sections"]
sentences_raw = draft_raw["sections"][0]["paragraphs"][0]["sentences"]
assert len(sentences_raw) == 3
assert all("uid" in s and "origin" in s and "citation" in s for s in sentences_raw)
assert sorted(draft_raw["done"]) == sorted(pending_uids)

draf_murni = draft_file.read_text(encoding="utf-8")

# Tandai per paragraf: pilih satu kalimat, tombol paragraf menandai semua kalimatnya.
report_tab._select_sentence(para_uid)
jumlah_kalimat_para = len(report_tab._selected_para_uids)
assert jumlah_kalimat_para >= 1, "Pemilihan paragraf kosong"
report_tab._mark_paragraph_done()
assert set(report_tab._selected_para_uids) <= report_tab._done, (
    "Tandai paragraf tidak menandai seluruh kalimatnya"
)

# Hapus sumber: sitasi merujuknya dilepas, indeks sumber lain digeser.
import tkinter.messagebox as _mb5

_mb5.askyesno = lambda *a, **k: True
jumlah_awal = len(report_tab._sources)
report_tab._source_listbox.selection_clear(0, "end")
report_tab._source_listbox.selection_set(0)
report_tab._hapus_source()
root3.update()
assert len(report_tab._sources) == jumlah_awal - 1, "Sumber tidak terhapus"
assert report_tab._sentence_by_uid(para_uid).source_ids == [0], (
    "Indeks sumber setelah terhapus tidak digeser"
)
tanpa_sumber = [s for s in report_tab._iter_sentences() if not s.source_ids]
assert all(not s.citation for s in tanpa_sumber), "Sitasi sumber yang dihapus harus dilepas"

# Kembalikan draf ke keadaan sebelum uji paragraf/hapus agar pemulihan root4 murni.
draft_file.write_text(draf_murni, encoding="utf-8")

root3.destroy()

# 6) Draf dipulihkan di jendela baru: sumber, judul, outline, template.
root4 = app.build_app()
assert root4 is not None
root4.update()
nb4 = find_notebook(root4)
tab_ids4 = nb4.tabs()
nb4.select(tab_ids4[2])
root4.update()
report_tab4 = nb4.nametowidget(tab_ids4[2]).controller
assert len(report_tab4._sources) == 2, "Sumber dari draf tidak dipulihkan"
assert report_tab4._sources[0]["title"] == "Judul Koreksi"
assert report_tab4._judul_var.get() == "Pertanian Organik"
assert report_tab4._outline_content() == ["Pendahuluan", "Tinjauan Pustaka"]
assert report_tab4._template is not None and report_tab4._template.nama == "Template Final"
assert len(report_tab4._sections) == 1, "Sections dari draf tidak dipulihkan"
assert len(report_tab4._statuses) == 3, "Status tidak dihitung ulang setelah pulih"
assert any(sentence.text == "Parafrase uji." for sentence in report_tab4._iter_sentences()), (
    "Hasil parafrase tidak dipulihkan dari draf"
)
assert not any("diedit manual" in sentence.text for sentence in report_tab4._iter_sentences()), (
    "Kalimat lama sebelum parafrase masih ada"
)
assert report_module.can_export(
    report_tab4._statuses, frozenset(report_tab4._done)
) == (True, ""), "Tanda selesai tidak dipulihkan dari draf"

# 7) Ekspor: gerbang memblokir saat tanda selesai dikosongkan; izinkan -> DOCX nyata + draf terhapus.
report_tab4._picker_vars["dosen"].set("Siti Aminah")
report_tab4._picker_vars["matkul"].set("Dasar Pertanian")

askyesno_flag = {"value": True}
tab_template_module.messagebox.askyesno = lambda *a, **k: askyesno_flag["value"]
tab_template_module.messagebox.showinfo = lambda *a, **k: dialogs.setdefault("info", []).append(a)

report_tab4._done.clear()
askyesno_flag["value"] = False
out_path = paths.OUTPUT_DIR / "Pertanian_Organik.docx"
out_path.unlink(missing_ok=True)
report_tab4._save_docx()
root4.update()
assert not out_path.exists(), "Ekspor tidak boleh jalan saat user menolak konfirmasi gerbang"

askyesno_flag["value"] = True
report_tab4._save_docx()
assert pump(root4, lambda: out_path.exists() and not draft_file.exists(), timeout_s=25.0), (
    "Render DOCX tidak menghasilkan berkas atau draf tidak terhapus; "
    f"status='{report_tab4._status.cget('text')}' "
    f"errors={[str(e) for e in dialogs['error']]} "
    f"info={dialogs.get('info')} file_exists={out_path.exists()}"
)
exported = Document(out_path)
text = "\n".join(par.text for par in exported.paragraphs)
assert "{{" not in text, [par.text for par in exported.paragraphs if "{{" in par.text]
assert "Pertanian Organik" in text, "Judul tidak terisi di dokumen"
assert "Laporan Pertanian" in text, "Referensi sumber yang dikutip tidak tercetak"

# 8) Generate ulang: catatan manual dipertahankan bila user menolak timpa; dibuang bila setuju.
manual_target = report_tab4._sections[0].paragraphs[0].sentences[0]
report_tab4._sections[0].paragraphs[0].sentences[0] = report_module.apply_manual_edit(
    manual_target, "Catatan manual saya."
)
regen_section = Section(
    2, "Bab Baru", [Paragraph([Sentence("Kalimat baru.", "q", [0], is_factual=False)])]
)
tab_report_module.content.generate_sections = lambda *args, **kwargs: [regen_section]


def _all_texts():
    return [s.text for s in report_tab4._iter_sentences()]


askyesno_flag["value"] = False  # Tidak = pertahankan catatan manual
report_tab4._generate()
assert pump(root4, lambda: any(sec.title == "Bab Baru" for sec in report_tab4._sections)), (
    "Generate ulang tidak selesai"
)
assert "Catatan manual saya." in _all_texts() and "Kalimat baru." in _all_texts(), (
    "Catatan manual harus dipertahankan saat user menolak tulis ulang semua"
)

askyesno_flag["value"] = True  # Ya = tulis ulang semua
report_tab4._generate()
assert pump(root4, lambda: "Catatan manual saya." not in _all_texts()), (
    "Catatan manual harus dibuang saat user memilih tulis ulang semua"
)
assert any(sec.title == "Bab Baru" for sec in report_tab4._sections)

# 9) Tab Data: Ubah dan Hapus bekerja pada baris terpilih (bug klik ditemukan saat uji manual).
data_inner = find_notebook(nb4.nametowidget(nb4.tabs()[0]))
dosen_panel = data_inner.nametowidget(data_inner.tabs()[1])  # sub-tab Dosen
dosen_panel._vars["nama"].set("Panel Uji Ubah")
dosen_panel._vars["gelar"].set("M.Si")
dosen_panel._vars["identitas"].set("000111")
dosen_panel._add()
assert any("Panel Uji Ubah" in t for t in dosen_panel._listbox.get(0, ttk_tk.END)), (
    "Dosen uji tidak masuk daftar"
)
idx = list(dosen_panel._listbox.get(0, ttk_tk.END)).index("Panel Uji Ubah (000111)")
dosen_panel._listbox.selection_set(idx)
dosen_panel._on_select()
dosen_panel._vars["nama"].set("Panel Uji Sudah Diubah")
dosen_panel._update()
root4.update()
assert any("Panel Uji Sudah Diubah" in t for t in dosen_panel._listbox.get(0, ttk_tk.END)), (
    "Tombol Ubah tidak mengubah baris terpilih"
)
idx = list(dosen_panel._listbox.get(0, ttk_tk.END)).index("Panel Uji Sudah Diubah (000111)")
dosen_panel._listbox.selection_set(idx)
dosen_panel._delete()
root4.update()
assert not any("Panel Uji" in t for t in dosen_panel._listbox.get(0, ttk_tk.END)), (
    "Tombol Hapus tidak menghapus baris terpilih"
)

# 10) Saat generate berjalan: tombol Generate terkunci, status memuat detik berjalan.
import time as _time
import tkinter.messagebox as _mb

_mb.askyesno = lambda *a, **k: True  # timpa isi laporan lama tanpa dialog

asli_generate = tab_report_module.content.generate_sections


def _generate_lambat(*a, **k):
    _time.sleep(1.2)
    return report_module.sections_from_dicts(
        [
            {
                "title": "Bab Kunci",
                "paragraphs": [
                    {
                        "sentences": [
                            {
                                "uid": "kunci1",
                                "text": "Kalimat uji tombol terkunci.",
                                "source_ids": [0],
                                "origin": "ai",
                                "citation": True,
                            }
                        ]
                    }
                ],
            }
        ]
    )


tab_report_module.content.generate_sections = _generate_lambat
try:
    report_tab4._generate()
    root4.update()
    assert str(report_tab4._generate_button.cget("state")) == "disabled", (
        "Tombol Generate tidak terkunci selama pekerjaan berjalan"
    )
    batas = _time.time() + 8
    while str(report_tab4._generate_button.cget("state")) != "normal":
        root4.update()
        _time.sleep(0.05)
        assert _time.time() < batas, "Tombol Generate tidak aktif kembali"
finally:
    tab_report_module.content.generate_sections = asli_generate

root4.update()
assert report_tab4._status.cget("text").startswith("Generate selesai: 1 bab."), (
    report_tab4._status.cget("text")
)
assert any(
    s.text == "Kalimat uji tombol terkunci." for s in report_tab4._iter_sentences()
), "Hasil generate palsu tidak masuk daftar kalimat"

root4.destroy()

print("SMOKE OK: jendela terbuka, 5 tab + 4 sub-tab ada, tambah dosen/matkul, tersimpan setelah dibuka ulang, Pengaturan menyimpan config.json, Template validasi+hapus jalan, Laporan sumber+metadata+draf jalan")
