"""Regresi temuan UAT 2026-10-24: persistensi draf, ekspor, verifikasi, dedupe.

Menutup temuan P0 hasil penggunaan end-to-end:
- save_active_draft dari tab Review (payload tanpa "fields") tidak boleh
  menghapus JUDUL/sumber yang sudah ada.
- clear_active_draft ("Mulai Lembar Kerja Kosong Baru") harus mengarsipkan
  draf yang masih berisi, bukan memusnahkannya.
- verify_document tanpa sumber dari pemanggil memakai sumber draf aktif
  (dulu: semua kalimat verbatim berubah merah setelah ganti tab).
- Ekspor setelah "restart" (jembatan baru, draf dari disk) tetap menghasilkan
  dokumen dengan sampul terisi, sitasi, dan tanpa penanda mentah.
"""

from pathlib import Path

import docx

from generator import database, paths
from generator.ai.content import Paragraph, Section, Sentence
from generator.web import api_bridge
from generator.web.api_bridge import ApiBridge


def _make_bridge(tmp_path: Path):
    paths.set_data_dir(tmp_path)
    paths.OUTPUT_DIR = tmp_path / "output"
    paths.OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    db = database.Database(tmp_path / "test.db")
    db.init_schema()
    bridge = ApiBridge(db, config_path=tmp_path / "config.json")
    return db, bridge


def _full_state(sources: list[dict], sections: list[dict]) -> dict:
    return {
        "fields": {
            "JUDUL": "Pengaruh Gamifikasi terhadap Motivasi Belajar",
            "DOSEN": "Dr. Ir. Soekarno Hadi, M.P.",
            "MATA_KULIAH": "Metodologi Penelitian",
            "TEMPLATE_ID": 1,
            "VARIANT": "B",
        },
        "sources": sources,
        "sections": sections,
        "done": [],
    }


def _source_record() -> dict:
    return {
        "authors": [{"family": "Sari", "given": "R. D."}],
        "year": "2023",
        "title": "Pengaruh Gamifikasi terhadap Motivasi Belajar",
        "container": "Jurnal Teknologi Pendidikan Indonesia",
        "text": (
            "Perkembangan teknologi informasi telah mengubah cara pembelajaran "
            "diselenggarakan di perguruan tinggi. Penelitian awal menunjukkan bahwa "
            "62 persen mahasiswa melaporkan penurunan motivasi belajar."
        ),
    }


def _section_dict(text: str, quote: str) -> dict:
    return {
        "level": 2,
        "title": "1.1 Latar Belakang",
        "paragraphs": [
            {
                "sentences": [
                    {
                        "text": text,
                        "quote": quote,
                        "source_ids": [0],
                        "is_factual": True,
                        "citation": True,
                    }
                ]
            }
        ],
    }


# --- P0-1: payload parsial tidak menghapus fields/sources ---------------------


def test_review_style_save_tidak_menghapus_fields_dan_sources(tmp_path: Path):
    db, bridge = _make_bridge(tmp_path)
    try:
        state = _full_state([_source_record()], [_section_dict("Kalimat pertama.", "Kalimat pertama.")])
        assert bridge.save_active_draft(state)["success"] is True

        # Persis seperti saveCurrentDraft milik TabReview: tanpa "fields",
        # sources kosong karena state komponen memang tidak memuatnya.
        review_payload = {
            "sections": state["sections"],
            "sources": [],
            "done": ["uid-1"],
        }
        assert bridge.save_active_draft(review_payload)["success"] is True

        reloaded = bridge.load_active_draft()
        assert reloaded["has_draft"] is True
        assert reloaded["fields"].get("JUDUL") == "Pengaruh Gamifikasi terhadap Motivasi Belajar"
        assert len(reloaded["sources"]) == 1
        assert reloaded["done"] == ["uid-1"]
    finally:
        db.close()


def test_autosave_studio_tanpa_berkas_draf_tidak_crash(tmp_path: Path):
    """Payload parsial Studio (fields/sources, tanpa sections/done) pada draf
    yang belum pernah ada tidak boleh melempar AttributeError (bug produksi:
    autosave gagal berulang setelah "Mulai Lembar Kerja Kosong Baru")."""
    db, bridge = _make_bridge(tmp_path)
    try:
        # Draf aktif belum pernah tersimpan -> existing = None
        result = bridge.save_active_draft({
            "fields": {"JUDUL": "Draf Baru", "VARIANT": "B"},
            "sources": [_source_record()],
        })
        assert result["success"] is True

        reloaded = bridge.load_active_draft()
        assert reloaded["has_draft"] is True
        assert reloaded["fields"].get("JUDUL") == "Draf Baru"
        assert reloaded["sections"] == []
        assert reloaded["done"] == []

        # Payload kosong total juga harus aman.
        assert bridge.save_active_draft({})["success"] is True
    finally:
        db.close()


def test_fields_eksplisit_dan_sumber_eksplisit_tetap_menang(tmp_path: Path):
    db, bridge = _make_bridge(tmp_path)
    try:
        state = _full_state([_source_record()], [_section_dict("Kalimat.", "Kalimat.")])
        bridge.save_active_draft(state)

        # TabLaporan mengirim fields lengkap + daftar sumber aktual: harapan
        # pemakaian normal adalah nilai baru menang (termasuk menghapus sumber).
        updated = {
            "fields": {"JUDUL": "Judul Baru Hasil Suntingan"},
            "sources": [],
            "sections": state["sections"],
            "done": [],
        }
        bridge.save_active_draft(updated)
        reloaded = bridge.load_active_draft()
        assert reloaded["fields"].get("JUDUL") == "Judul Baru Hasil Suntingan"
        # DOSEN dari simpanan sebelumnya tidak ikut terhapus (gabung per kunci)
        assert reloaded["fields"].get("DOSEN") == "Dr. Ir. Soekarno Hadi, M.P."
        assert reloaded["sources"] == []
    finally:
        db.close()


# --- P0-4: lembar kerja baru mengarsipkan draf berisi -------------------------


def test_clear_active_draft_mengarsipkan_draf_berisi(tmp_path: Path):
    db, bridge = _make_bridge(tmp_path)
    try:
        state = _full_state([_source_record()], [_section_dict("Kalimat.", "Kalimat.")])
        bridge.save_active_draft(state)

        assert bridge.clear_active_draft()["success"] is True

        assert bridge.load_active_draft()["has_draft"] is False
        arsip = [d for d in bridge.list_saved_drafts() if d["draft_id"].startswith("arsip_")]
        assert len(arsip) == 1
        assert arsip[0]["sentence_count"] == 1
        assert arsip[0]["sources_count"] == 1

        # Draf kosong boleh dihapus tanpa meninggalkan arsip
        bridge.clear_active_draft()
        arsip = [d for d in bridge.list_saved_drafts() if d["draft_id"].startswith("arsip_")]
        assert len(arsip) == 1
    finally:
        db.close()


# --- P0-3: verifikasi memakai sumber draf aktif & dedupe per bab --------------


def test_verify_document_jatuh_ke_sumber_draf_aktif(tmp_path: Path):
    db, bridge = _make_bridge(tmp_path)
    try:
        teks = (
            "Perkembangan teknologi informasi telah mengubah cara pembelajaran "
            "diselenggarakan di perguruan tinggi."
        )
        state = _full_state([_source_record()], [_section_dict(teks, teks)])
        bridge.save_active_draft(state)

        # Jembatan baru = simulasi restart aplikasi; pemanggil (tab Review)
        # mengirim daftar sumber kosong.
        bridge2 = ApiBridge(db, config_path=tmp_path / "config.json")
        statuses = bridge2.verify_document(state["sections"], [])
        assert len(statuses) == 1
        assert statuses[0]["status"] == "terverifikasi"
    finally:
        db.close()


def test_merge_sections_dedupe_hanya_dalam_bab_sama():
    from generator.ai.content import _merge_sections

    s_latar = Sentence(
        text="Gamifikasi meningkatkan keterlibatan mahasiswa dalam forum diskusi daring hingga 45 persen.",
        quote="keterlibatan mahasiswa dalam forum diskusi daring hingga 45 persen",
        source_ids=[0],
        is_factual=True,
        uid="u1",
    )
    s_hasil_mirip = Sentence(
        text="Hasil penelitian membuktikan gamifikasi meningkatkan keterlibatan mahasiswa dalam forum diskusi daring.",
        quote="keterlibatan mahasiswa dalam forum diskusi daring",
        source_ids=[0],
        is_factual=True,
        uid="u2",
    )
    s_duplikat_latar = Sentence(
        text="Gamifikasi meningkatkan keterlibatan mahasiswa dalam forum diskusi daring hingga 45 persen.",
        quote="sama",
        source_ids=[0],
        is_factual=True,
        uid="u3",
    )

    existing = [
        Section(1, "BAB 1. PENDAHULUAN", [Paragraph([s_latar])]),
        Section(1, "BAB 3. HASIL DAN PEMBAHASAN", []),
    ]
    incoming = [
        Section(1, "BAB 1. PENDAHULUAN", [Paragraph([s_duplikat_latar])]),
        Section(1, "BAB 3. HASIL DAN PEMBAHASAN", [Paragraph([s_hasil_mirip])]),
    ]
    _merge_sections(existing, incoming)

    bab1 = next(s for s in existing if "PENDAHULUAN" in s.title)
    bab3 = next(s for s in existing if "HASIL" in s.title)
    assert len(bab1.paragraphs) == 1  # duplikat dalam bab yang sama dibuang
    assert len(bab3.paragraphs) == 1  # kalimat mirip lintas bab tetap hidup


# --- P0-2: ekspor pasca-restart tetap benar -----------------------------------


def test_export_setelah_restart_lengkap_tanpa_penanda_mentah(tmp_path: Path, monkeypatch):
    from generator.web import api_bridge as api_bridge_module

    db, bridge = _make_bridge(tmp_path)
    try:
        bridge.save_profile(
            {"nama": "Muhammad Azhar Ghaelby", "nim": "232410101088", "prodi": "Sistem Informasi"}
        )
        tpl_res = bridge.create_default_template({"font_name": "Times New Roman", "margin_preset": "4433"})
        assert tpl_res.get("success") is True
        tpl_id = bridge.list_templates()[0]["id"]

        # Simpan draf TANPA NAMA/NIM/TANGGAL di fields (skenario nyata: sampul
        # diisi dari profil saat ekspor), lalu tutup aplikasi.
        state = _full_state([_source_record()], [_section_dict(
            "Perkembangan teknologi informasi telah mengubah cara pembelajaran diselenggarakan di perguruan tinggi.",
            "Perkembangan teknologi informasi telah mengubah cara pembelajaran diselenggarakan di perguruan tinggi.",
        )])
        bridge.save_active_draft(state)
        del bridge

        # Aplikasi dibuka kembali: jembatan segar dari disk.
        monkeypatch.setattr(api_bridge_module.os, "startfile", lambda p: None, raising=False)
        bridge2 = ApiBridge(db, config_path=tmp_path / "config.json")
        restored = bridge2.load_active_draft()
        assert restored["has_draft"] is True
        assert restored["sources"], "sumber harus ikut tersimpan di draf"

        export_state = dict(restored)
        export_state["fields"] = {**restored["fields"], "TEMPLATE_ID": tpl_id}
        export_state["template_id"] = tpl_id
        export_state["open_word"] = False
        res = bridge2.export_report(export_state)
        assert res.get("success") is True, res.get("error")

        doc = docx.Document(res["output_path"])
        doc_text = "\n".join(p.text for p in doc.paragraphs)
        assert "{{" not in doc_text
        assert "Muhammad Azhar Ghaelby" in doc_text
        assert "232410101088" in doc_text
        assert "Pengaruh Gamifikasi terhadap Motivasi Belajar" in doc_text
    finally:
        db.close()


# --- P0-4: preflight menandai kalimat faktual tanpa sitasi & bagian kosong ----


def test_preflight_menandai_faktual_tanpa_sitasi_dan_bagian_kosong(tmp_path: Path):
    from generator import preflight
    from generator.ai.citation import Author, Source
    from generator.report import sections_from_dicts
    from generator.verification import verify_sections

    source = Source(authors=[Author(family="Sari", given="R.")], year="2023", title="Gamifikasi")
    sections = sections_from_dicts([
        {
            "level": 1,
            "title": "BAB 2. TINJAUAN PUSTAKA",
            "paragraphs": [],
        },
        {
            "level": 2,
            "title": "2.1 Latar Belakang",
            "paragraphs": [
                {
                    "sentences": [
                        {
                            "text": "Klaim tanpa sitasi sama sekali.",
                            "quote": "klaim",
                            "source_ids": [],
                            "is_factual": True,
                            "citation": True,
                        }
                    ]
                }
            ],
        },
        {
            "level": 1,
            "title": "BAB 3. METODOLOGI",
            "paragraphs": [],
        },
    ])
    statuses = verify_sections(sections, ["klaim"])

    issues = preflight.find_issues(
        judul="Laporan Uji",
        sources=[source],
        statuses=statuses,
        cited=set(),
        logo_path=None,
        sections=sections,
    )
    kinds = {issue.kind for issue in issues}
    assert "citations" in kinds
    # BAB 2 punya isi lewat subbab 2.1, jadi tidak dihitung kosong (F6);
    # BAB 3 benar-benar tanpa isi di seluruh subbabnya, tetap ditandai.
    empty_issue = next(issue for issue in issues if issue.kind == "empty_sections")
    assert "BAB 3. METODOLOGI" in empty_issue.message
    assert "BAB 2. TINJAUAN PUSTAKA" not in empty_issue.message

    checks = preflight.run_checks(
        judul="Laporan Uji",
        sources=[source],
        statuses=statuses,
        cited=set(),
        logo_path=None,
        sections=sections,
    )
    labels = {c.label: c for c in checks}
    assert labels["Sitasi & pustaka konsisten"].ok is False
    assert "Bagian tanpa isi" in labels
    assert labels["Bagian tanpa isi"].ok is False
