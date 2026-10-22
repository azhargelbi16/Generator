from docx import Document
from generator.document import placeholders


def test_replaces_split_runs_and_keeps_neighbours():
    doc = Document()
    p = doc.add_paragraph()
    p.add_run("A {{JU")
    p.add_run("DUL}} B")
    replaced = placeholders.replace_in_paragraph(p, {"{{JUDUL}}": "X"})
    assert replaced == ["{{JUDUL}}"]
    assert p.text == "A X B"


def test_multiple_placeholders_one_paragraph():
    doc = Document()
    p = doc.add_paragraph()
    p.add_run("{{NAMA}} dan {{NIM}}")
    placeholders.replace_in_paragraph(p, {"{{NAMA}}": "Budi", "{{NIM}}": "1"})
    assert p.text == "Budi dan 1"


def test_is_standalone():
    doc = Document()
    a = doc.add_paragraph("{{LOGO}}")
    a.add_run("")
    b = doc.add_paragraph("lihat {{LOGO}} ini")
    assert placeholders.is_standalone(a, "{{LOGO}}") is True
    assert placeholders.is_standalone(b, "{{LOGO}}") is False


def test_collect_names_includes_table_cells():
    doc = Document()
    t = doc.add_table(rows=1, cols=1)
    t.cell(0, 0).paragraphs[0].add_run("{{DOSEN}}")
    assert "{{DOSEN}}" in placeholders.collect_names(doc)


def _add_hyperlink(paragraph, text):
    from docx.opc.constants import RELATIONSHIP_TYPE as RT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    rel_id = paragraph.part.relate_to("https://example.com", RT.HYPERLINK, is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), rel_id)
    run = OxmlElement("w:r")
    node = OxmlElement("w:t")
    node.text = text
    run.append(node)
    link.append(run)
    paragraph._p.append(link)


def test_penanda_di_dalam_hyperlink_ikut_diganti():
    doc = Document()
    p = doc.add_paragraph("Lihat ")
    _add_hyperlink(p, "{{JUDUL}}")
    assert "{{JUDUL}}" in placeholders.collect_names(doc)

    replaced = placeholders.replace_in_paragraph(p, {"{{JUDUL}}": "Laporan X"})
    assert replaced == ["{{JUDUL}}"]
    assert p.text == "Lihat Laporan X"


def test_collect_names_mencakup_header_pertama_dan_halaman_genap():
    doc = Document()
    section = doc.sections[0]
    section.different_first_page_header_footer = True
    section.first_page_header.paragraphs[0].add_run("{{JUDUL}}")
    doc.settings.odd_and_even_pages_header_footer = True
    section.even_page_footer.paragraphs[0].add_run("{{NIM}}")

    names = placeholders.collect_names(doc)
    assert "{{JUDUL}}" in names
    assert "{{NIM}}" in names
