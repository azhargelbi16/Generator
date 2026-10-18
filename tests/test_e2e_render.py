"""Uji end-to-end render: template penuh menghasilkan dokumen bersih (Task 18)."""

from docx import Document

from generator.ai.content import Paragraph, Section, Sentence
from generator.document import renderer, template


def test_full_template_produces_clean_report(tmp_path):
    p = tmp_path / "t.docx"
    d = Document()
    for ph in template.REQUIRED_PLACEHOLDERS:
        d.add_paragraph(ph)
    d.save(p)
    secs = [Section(1, "Bab Satu", [Paragraph([Sentence("Isi.", "x", [0])])])]
    out = renderer.render_report(
        p,
        {"{{JUDUL}}": "Laporan", "{{NAMA}}": "Budi", "{{NIM}}": "1",
         "{{MATA_KULIAH}}": "Agro", "{{DOSEN}}": "Dr. X", "{{TANGGAL}}": "2026"},
        secs, ["Ref."], None, tmp_path / "o.docx",
    )
    text = "\n".join(par.text for par in Document(out).paragraphs)
    assert "{{" not in text and "Laporan" in text and "Ref." in text and "Bab Satu" in text
