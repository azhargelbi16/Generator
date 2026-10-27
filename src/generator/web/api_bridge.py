"""API Bridge connecting pywebview frontend to Generator Laporan core logic."""

from __future__ import annotations

import base64
import concurrent.futures
import json
import logging
import os
import re
import shutil
import threading
from dataclasses import asdict
from datetime import datetime
from pathlib import Path
from typing import Any, Iterator

import webview

from .. import atomic, config, database, drafts, format_spec, paths, preflight, report
from ..ai import bibliography, content, citation, client, dosen
from ..ai.citation import Source, Author
from ..ai.client import AIClient, AIError
from ..document import renderer, template, placeholders
from ..sources import pdf_reader, web_reader, doi_resolver
from ..sources import intake
from ..sources.records import (
    DRAFT_ACTIVE_ID,
    normalize_metadata,
    record_to_source,
    source_to_record,
)
from .. import reviewed
from ..verification import verify_sections, STATUS_VERIFIED, STATUS_NEEDS_REVIEW, STATUS_UNVERIFIED, STATUS_NON_FACTUAL

log = logging.getLogger("generator.web.api")

_PLACEHOLDER_RE = re.compile(r"\{\{\s*([A-Z0-9_]+)\s*\}\}")


def _leftover_placeholders(docx_path: Path, limit: int = 5) -> list[str]:
    """Penanda {{...}} yang masih tersisa di dokumen hasil render.

    Memeriksa paragraf, tabel, serta header/footer; janji aplikasi adalah
    penanda tidak pernah tercetak mentah di berkas keluaran.
    """
    from docx import Document  # impor lokal: python-docx cukup berat untuk startup

    try:
        document = Document(str(docx_path))
    except Exception:
        log.exception("Gagal membuka hasil ekspor untuk pemeriksaan penanda")
        return []

    def _texts() -> Iterator[str]:
        for paragraph in document.paragraphs:
            yield paragraph.text
        for table in document.tables:
            for row in table.rows:
                for cell in row.cells:
                    for paragraph in cell.paragraphs:
                        yield paragraph.text
        for section in document.sections:
            for part in (section.header, section.footer):
                for paragraph in part.paragraphs:
                    yield paragraph.text

    found: list[str] = []
    for text in _texts():
        for match in _PLACEHOLDER_RE.finditer(text):
            marker = match.group(1)
            if marker not in found:
                found.append(marker)
                if len(found) >= limit:
                    return found
    return found


class ApiBridge:
    """Methods exposed to JavaScript via window.pywebview.api.*"""

    def __init__(self, db: database.Database, config_path: Path | None = None):
        self.db = db
        self.config_path = Path(config_path) if config_path else Path(paths.CONFIG_PATH)
        self._cfg = config.load_config(self.config_path)
        self._window: webview.Window | None = None
        self._cancel_gen_event = threading.Event()
        self._gen_thread: threading.Thread | None = None
        self._enrichment_semaphore = threading.Semaphore(2)

    def set_window(self, window: webview.Window) -> None:
        self._window = window

    def _emit(self, event_name: str, payload: Any = None) -> None:
        """Emit custom DOM event to the frontend."""
        if not self._window:
            return
        data_json = json.dumps(payload or {})
        script = f"window.dispatchEvent(new CustomEvent('{event_name}', {{detail: {data_json}}}));"
        try:
            self._window.evaluate_js(script)
        except Exception as e:
            log.warning(f"Failed to emit JS event {event_name}: {e}")

    # =========================================================================
    # VERSIONING & API CONTRACT (Formal Schema v2.0.0)
    # =========================================================================

    def get_api_version(self) -> str:
        from .contract import API_VERSION
        return API_VERSION

    def get_api_contract(self) -> dict[str, Any]:
        from .contract import get_contract_metadata
        return get_contract_metadata()

    # =========================================================================
    # CONFIG & AI SERVICE
    # =========================================================================

    def get_config(self) -> dict[str, Any]:
        self._cfg = config.load_config(self.config_path)
        data = asdict(self._cfg)
        data["ai_api_key_set"] = bool(self._cfg.ai_api_key and self._cfg.ai_api_key.strip())
        data["ai_api_key"] = ""
        return data

    def save_config(self, cfg_data: dict[str, Any]) -> dict[str, Any]:
        try:
            temp = float(str(cfg_data.get("ai_temperature", 0.2)).replace(",", "."))
            temp = max(0.0, min(2.0, temp))
        except (ValueError, TypeError):
            temp = 0.2

        new_key = str(cfg_data.get("ai_api_key", "")).strip()
        if not new_key:
            api_key = self._cfg.ai_api_key
        else:
            api_key = new_key

        self._cfg = config.Config(
            ai_base_url=str(cfg_data.get("ai_base_url", "http://localhost:1234/v1")).strip(),
            ai_api_key=api_key,
            ai_model=str(cfg_data.get("ai_model", "gpt-4o-mini")).strip(),
            ai_temperature=temp,
            citation_style=str(cfg_data.get("citation_style", "et_al")),
            dosen_review=bool(cfg_data.get("dosen_review", False)),
        )
        config.save_config(self._cfg, self.config_path)
        res_data = asdict(self._cfg)
        res_data["ai_api_key_set"] = bool(self._cfg.ai_api_key and self._cfg.ai_api_key.strip())
        res_data["ai_api_key"] = ""
        return {"success": True, "config": res_data}

    def test_ai_connection(self, base_url: str | None = None, api_key: str | None = None, model: str | None = None) -> dict[str, Any]:
        key_to_use = str(api_key).strip() if (api_key and str(api_key).strip()) else self._cfg.ai_api_key
        test_cfg = config.Config(
            ai_base_url=str(base_url if base_url else self._cfg.ai_base_url).strip(),
            ai_api_key=key_to_use,
            ai_model=str(model if model is not None else self._cfg.ai_model).strip(),
            ai_temperature=self._cfg.ai_temperature,
            citation_style=self._cfg.citation_style,
            dosen_review=self._cfg.dosen_review,
        )
        cl = AIClient(test_cfg)
        ok, msg = client.check_connection(cl)
        return {"ok": ok, "message": msg}

    def fetch_models(self, base_url: str | None = None, api_key: str | None = None) -> dict[str, Any]:
        key_to_use = str(api_key).strip() if (api_key and str(api_key).strip()) else self._cfg.ai_api_key
        test_cfg = config.Config(
            ai_base_url=str(base_url if base_url else self._cfg.ai_base_url).strip(),
            ai_api_key=key_to_use,
            ai_model=self._cfg.ai_model,
            ai_temperature=self._cfg.ai_temperature,
            citation_style=self._cfg.citation_style,
            dosen_review=self._cfg.dosen_review,
        )
        try:
            models = client.list_models(test_cfg)
            return {"ok": True, "models": models}
        except Exception as e:
            return {"ok": False, "error": str(e), "models": []}

    # =========================================================================
    # DATABASE & ENTITIES
    # =========================================================================

    def get_profile(self) -> dict[str, str]:
        p = self.db.get_profil()
        return {
            "nama": p.nama,
            "nim": p.nim,
            "prodi": p.prodi,
            "fakultas": getattr(p, "fakultas", ""),
            "institusi": getattr(p, "institusi", ""),
        }

    def save_profile(self, data: dict[str, str]) -> dict[str, Any]:
        nama = str(data.get("nama", "")).strip()
        nim = str(data.get("nim", "")).strip()
        prodi = str(data.get("prodi", "")).strip()
        fakultas = str(data.get("fakultas", "")).strip()
        institusi = str(data.get("institusi", "")).strip()
        self.db.set_profil(
            database.Profil(
                nama=nama,
                nim=nim,
                prodi=prodi,
                fakultas=fakultas,
                institusi=institusi,
            )
        )
        try:
            active_draft = drafts.load_draft(DRAFT_ACTIVE_ID)
            if active_draft:
                f = dict(active_draft.fields) if active_draft.fields else {}
                f["NAMA"] = nama
                f["NIM"] = nim
                f["PRODI"] = prodi
                f["FAKULTAS"] = fakultas
                f["INSTITUSI"] = institusi
                drafts.save_draft(drafts.DraftState(
                    draft_id=DRAFT_ACTIVE_ID,
                    fields=f,
                    sources=active_draft.sources,
                    sections=active_draft.sections,
                    done=active_draft.done,
                ))
        except Exception:
            log.exception("Gagal menyinkronkan profil ke draf aktif")
        return {"success": True}

    def list_dosen(self) -> list[dict[str, Any]]:
        return [asdict(d) for d in self.db.list_dosen()]

    def add_dosen(self, d: dict[str, Any]) -> dict[str, Any]:
        dosen_obj = database.Dosen(
            id=0,
            nama=str(d.get("nama", "")).strip(),
            gelar=str(d.get("gelar", "")).strip(),
            identitas=str(d.get("identitas", "")).strip(),
        )
        new_id = self.db.add_dosen(dosen_obj)
        return {"success": True, "dosen": {"id": new_id, "nama": dosen_obj.nama, "gelar": dosen_obj.gelar, "identitas": dosen_obj.identitas}}

    def update_dosen(self, d: dict[str, Any]) -> dict[str, Any]:
        dosen_obj = database.Dosen(
            id=int(d["id"]),
            nama=str(d.get("nama", "")).strip(),
            gelar=str(d.get("gelar", "")).strip(),
            identitas=str(d.get("identitas", "")).strip(),
        )
        self.db.update_dosen(dosen_obj)
        return {"success": True}

    def delete_dosen(self, dosen_id: int) -> dict[str, Any]:
        self.db.delete_dosen(int(dosen_id))
        return {"success": True}

    def list_matkul(self) -> list[dict[str, Any]]:
        return [asdict(m) for m in self.db.list_mata_kuliah()]

    def add_matkul(self, m: dict[str, Any]) -> dict[str, Any]:
        matkul_obj = database.MataKuliah(
            id=0,
            nama=str(m.get("nama", "")).strip(),
            kode=str(m.get("kode", "")).strip(),
        )
        new_id = self.db.add_mata_kuliah(matkul_obj)
        return {"success": True, "matkul": {"id": new_id, "nama": matkul_obj.nama, "kode": matkul_obj.kode}}

    def update_matkul(self, m: dict[str, Any]) -> dict[str, Any]:
        matkul_obj = database.MataKuliah(
            id=int(m["id"]),
            nama=str(m.get("nama", "")).strip(),
            kode=str(m.get("kode", "")).strip(),
        )
        self.db.update_mata_kuliah(matkul_obj)
        return {"success": True}

    def delete_matkul(self, mk_id: int) -> dict[str, Any]:
        self.db.delete_mata_kuliah(int(mk_id))
        return {"success": True}

    def _read_image_as_data_url(self, file_path: Path | str) -> str | None:
        try:
            p = Path(file_path)
            if not p.is_file():
                return None
            ext = p.suffix.lower()
            mime_map = {
                ".png": "image/png",
                ".jpg": "image/jpeg",
                ".jpeg": "image/jpeg",
                ".bmp": "image/bmp",
                ".webp": "image/webp",
                ".gif": "image/gif",
                ".svg": "image/svg+xml",
            }
            mime = mime_map.get(ext, "image/png")
            raw = p.read_bytes()
            encoded = base64.b64encode(raw).decode("ascii")
            return f"data:{mime};base64,{encoded}"
        except Exception:
            return None

    def list_logos(self) -> list[dict[str, Any]]:
        logos = self.db.list_logo()
        res = []
        paths.LOGOS_DIR.mkdir(parents=True, exist_ok=True)
        for l in logos:
            src_file = Path(l.path)
            # Jika logo masih berada di luar folder program (data/logos),
            # salin otomatis ke data/logos agar terorganisir rapi dan aman.
            if src_file.is_file():
                try:
                    src_file.resolve().relative_to(paths.LOGOS_DIR.resolve())
                except ValueError:
                    dest_file = paths.LOGOS_DIR / src_file.name
                    counter = 1
                    while dest_file.exists() and dest_file.resolve() != src_file.resolve():
                        dest_file = paths.LOGOS_DIR / f"{src_file.stem}_{counter}{src_file.suffix}"
                        counter += 1
                    if not dest_file.exists():
                        shutil.copy2(src_file, dest_file)
                    self.db.update_logo(database.Logo(l.id, l.nama, str(dest_file)))
                    l = database.Logo(l.id, l.nama, str(dest_file))
                    src_file = dest_file

            exists = src_file.is_file()
            data_url = self._read_image_as_data_url(src_file) if exists else None
            res.append({
                "id": l.id,
                "nama": l.nama,
                "path": l.path,
                "exists": exists,
                "data_url": data_url,
            })
        return res

    def upload_logo(self) -> dict[str, Any]:
        if not self._window:
            return {"cancelled": True}
        file_types = ("Image files (*.png;*.jpg;*.jpeg;*.bmp)", "All files (*.*)")
        res = self._window.create_file_dialog(webview.FileDialog.OPEN, allow_multiple=False, file_types=file_types)
        if not res or len(res) == 0:
            return {"cancelled": True}

        src_path = Path(res[0])
        paths.LOGOS_DIR.mkdir(parents=True, exist_ok=True)
        dest_path = paths.LOGOS_DIR / src_path.name
        counter = 1
        while dest_path.exists():
            dest_path = paths.LOGOS_DIR / f"{src_path.stem}_{counter}{src_path.suffix}"
            counter += 1

        shutil.copy2(src_path, dest_path)
        logo_id = self.db.add_logo(database.Logo(0, src_path.stem, str(dest_path)))
        data_url = self._read_image_as_data_url(dest_path)
        return {
            "success": True,
            "logo": {
                "id": logo_id,
                "nama": src_path.stem,
                "path": str(dest_path),
                "exists": True,
                "data_url": data_url,
            }
        }

    def open_logos_folder(self) -> dict[str, Any]:
        try:
            paths.LOGOS_DIR.mkdir(parents=True, exist_ok=True)
            os.startfile(paths.LOGOS_DIR)
            return {"success": True}
        except Exception as e:
            return {"error": str(e)}

    def delete_logo(self, logo_id: int) -> dict[str, Any]:
        logos = self.db.list_logo()
        found = next((l for l in logos if l.id == int(logo_id)), None)
        if found and Path(found.path).is_file():
            try:
                Path(found.path).unlink(missing_ok=True)
            except Exception:
                pass
        self.db.delete_logo(int(logo_id))
        return {"success": True}

    # =========================================================================
    # TEMPLATES
    # =========================================================================

    def _create_academic_template_docx(self, dest_path: Path, options: dict[str, Any] | None = None) -> None:
        from docx import Document
        from docx.shared import Cm, Pt, RGBColor
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.enum.style import WD_STYLE_TYPE

        opts = options or {}
        font_name = opts.get("font_name", "Times New Roman")
        margin_preset = opts.get("margin_preset", "4433")
        line_spacing = float(opts.get("line_spacing", 1.5))

        doc = Document()

        # 1. Dimensi Kertas A4 (21.0 x 29.7 cm) & Margin Baku
        section = doc.sections[0]
        section.page_width = Cm(21.0)
        section.page_height = Cm(29.7)
        if margin_preset == "4433":
            # Standar Skripsi/Tesis Dikti: Kiri 4 cm, Atas 4 cm, Kanan 3 cm, Bawah 3 cm
            section.top_margin = Cm(4.0)
            section.left_margin = Cm(4.0)
            section.right_margin = Cm(3.0)
            section.bottom_margin = Cm(3.0)
        else:
            # Standar Makalah/Artikel Bebas: 3 cm sekeliling
            section.top_margin = Cm(3.0)
            section.left_margin = Cm(3.0)
            section.right_margin = Cm(3.0)
            section.bottom_margin = Cm(3.0)

        # 2. Tipografi Baku
        normal = doc.styles["Normal"]
        normal.font.name = font_name
        normal.font.size = Pt(12)
        normal.font.color.rgb = RGBColor(0, 0, 0)
        normal.paragraph_format.line_spacing = line_spacing
        normal.paragraph_format.space_after = Pt(0)

        if "Heading 1" in doc.styles:
            h1 = doc.styles["Heading 1"]
        else:
            h1 = doc.styles.add_style("Heading 1", WD_STYLE_TYPE.PARAGRAPH)
        h1.font.name = font_name
        h1.font.size = Pt(14)
        h1.font.bold = True
        h1.font.color.rgb = RGBColor(0, 0, 0)
        h1.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        h1.paragraph_format.space_before = Pt(12)
        h1.paragraph_format.space_after = Pt(6)
        h1.paragraph_format.line_spacing = line_spacing

        if "Heading 2" in doc.styles:
            h2 = doc.styles["Heading 2"]
        else:
            h2 = doc.styles.add_style("Heading 2", WD_STYLE_TYPE.PARAGRAPH)
        h2.font.name = font_name
        h2.font.size = Pt(12)
        h2.font.bold = True
        h2.font.color.rgb = RGBColor(0, 0, 0)
        h2.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
        h2.paragraph_format.space_before = Pt(6)
        h2.paragraph_format.space_after = Pt(4)
        h2.paragraph_format.line_spacing = line_spacing

        # 3. Halaman Sampul (Cover Page) — struktur formal laporan akademik
        def _cover_line(
            text: str,
            *,
            bold: bool = False,
            size: int = 12,
            space_before: Any = None,
            space_after: Any = None,
            line_spacing: float = 1.15,
        ):
            """Satu baris sampul tersendiri, terpusat, dengan tipografi seragam."""
            par = doc.add_paragraph()
            fmt = par.paragraph_format
            fmt.alignment = WD_ALIGN_PARAGRAPH.CENTER
            fmt.line_spacing = line_spacing
            fmt.space_before = space_before if space_before is not None else Pt(0)
            fmt.space_after = space_after if space_after is not None else Pt(0)
            run = par.add_run(text)
            run.bold = bold
            run.font.name = font_name
            run.font.size = Pt(size)
            return par

        # Judul laporan: kapital, tebal, paling atas.
        _cover_line(
            "{{JUDUL}}", bold=True, size=16,
            space_before=Cm(1.0), space_after=Cm(0.8),
        )
        # Keterangan jenis dokumen.
        _cover_line("LAPORAN AKADEMIK", bold=True, space_after=Cm(1.0))
        # Logo institusi 5x5 cm; paragrafnya dihapus renderer bila logo kosong.
        _cover_line(
            "{{LOGO}}", space_before=Cm(0.5), space_after=Cm(1.2),
            line_spacing=1.0,
        )

        # Blok penyusun (tengah sampul).
        _cover_line("Disusun Oleh:", bold=True, space_before=Cm(0.5), space_after=Pt(6))
        _cover_line("{{NAMA}}", bold=True)
        _cover_line("NIM: {{NIM}}")
        _cover_line("Mata Kuliah: {{MATA_KULIAH}}", space_before=Pt(6))

        # Blok dosen pembimbing.
        _cover_line("Dosen Pengampu:", bold=True, space_before=Cm(0.8), space_after=Pt(6))
        _cover_line("{{DOSEN}}")
        _cover_line("{{TANGGAL}}", space_before=Pt(6))

        # Blok institusi formal di bagian bawah sampul.
        _cover_line("PROGRAM STUDI {{PRODI}}", space_before=Cm(2.5))
        _cover_line("FAKULTAS {{FAKULTAS}}")
        _cover_line("{{INSTITUSI}}", bold=True)
        _cover_line("{{TAHUN}}", space_before=Pt(6))

        doc.add_page_break()

        # 4. Halaman Daftar Isi
        h_toc = doc.add_paragraph(style="Heading 1")
        r_toc = h_toc.add_run("DAFTAR ISI")
        r_toc.bold = True
        h_toc.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER

        p_toc = doc.add_paragraph("{{DAFTAR_ISI}}")
        p_toc.paragraph_format.space_before = Pt(12)

        doc.add_page_break()

        # 5. Halaman Isi Laporan
        doc.add_paragraph("{{ISI_LAPORAN}}")

        doc.add_page_break()

        # 6. Halaman Daftar Pustaka
        h_pus = doc.add_paragraph(style="Heading 1")
        r_pus = h_pus.add_run("DAFTAR PUSTAKA")
        r_pus.bold = True
        h_pus.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER

        p_pus = doc.add_paragraph("{{DAFTAR_PUSTAKA}}")
        p_pus.paragraph_format.space_before = Pt(12)

        doc.save(dest_path)

    def _ensure_default_template(self) -> None:
        try:
            paths.TEMPLATES_DIR.mkdir(parents=True, exist_ok=True)
            dest_path = paths.TEMPLATES_DIR / "Template_Standar_Akademik.docx"
            if not dest_path.is_file():
                self._create_academic_template_docx(dest_path)

            all_tpls = self.db.list_template()
            desc = "Template Word resmi baku (Margin 4-4-3-3 cm, Times New Roman 12pt, Cover, TOC, APA 7)"
            if not any(Path(t.path).resolve() == dest_path.resolve() for t in all_tpls):
                self.db.add_template(database.Template(
                    0,
                    "Template Standar Akademik",
                    desc,
                    str(dest_path)
                ))
        except Exception as e:
            log.warning(f"Gagal memastikan template default: {e}")

    def create_default_template(self, options: dict[str, Any] | None = None) -> dict[str, Any]:
        try:
            paths.TEMPLATES_DIR.mkdir(parents=True, exist_ok=True)
            dest_path = paths.TEMPLATES_DIR / "Template_Standar_Akademik.docx"
            self._create_academic_template_docx(dest_path, options)

            all_tpls = self.db.list_template()
            existing = next((t for t in all_tpls if Path(t.path).resolve() == dest_path.resolve()), None)
            desc = "Template Word resmi baku (Margin 4-4-3-3 cm, Times New Roman 12pt, Cover, TOC, APA 7)"
            if existing:
                self.db.update_template(database.Template(
                    existing.id,
                    "Template Standar Akademik",
                    desc,
                    str(dest_path)
                ))
            else:
                self.db.add_template(database.Template(
                    0,
                    "Template Standar Akademik",
                    desc,
                    str(dest_path)
                ))
            return {"success": True, "path": str(dest_path)}
        except Exception as e:
            log.exception("Gagal membuat template standar")
            return {"success": False, "error": str(e)}

    def open_template(self, template_id: int) -> dict[str, Any]:
        tpl = next((t for t in self.db.list_template() if t.id == int(template_id)), None)
        if not tpl:
            return {"success": False, "error": "Template tidak ditemukan"}
        p = Path(tpl.path)
        if not p.is_file():
            return {"success": False, "error": f"Berkas tidak ditemukan: {p}"}
        if os.name == "nt":
            try:
                os.startfile(str(p.resolve()))
                return {"success": True}
            except Exception as e:
                return {"success": False, "error": f"Gagal membuka aplikasi Word: {e}"}
        return {"success": True}

    def open_templates_folder(self) -> dict[str, Any]:
        try:
            paths.TEMPLATES_DIR.mkdir(parents=True, exist_ok=True)
            if os.name == "nt":
                os.startfile(str(paths.TEMPLATES_DIR.resolve()))
            return {"success": True}
        except Exception as e:
            return {"success": False, "error": f"Gagal membuka folder template: {e}"}

    def list_templates(self) -> list[dict[str, Any]]:
        templates = self.db.list_template()
        if not templates:
            self._ensure_default_template()
            templates = self.db.list_template()
        results = []
        for t in templates:
            p = Path(t.path)
            status = "berkas_hilang"
            missing: list[str] = []
            not_standalone: list[str] = []
            unknown: list[str] = []
            placeholders_list: list[str] = []

            if p.is_file():
                try:
                    res = template.validate_template(p)
                    missing = list(res.missing)
                    not_standalone = list(res.not_standalone)
                    unknown = list(res.unknown)
                    status = "lengkap" if res.is_processable else "tidak_lengkap"
                    doc = template.load_template(p)
                    placeholders_list = placeholders.collect_names(doc)
                except Exception as e:
                    status = f"rusak: {e}"

            is_default = "Template_Standar_Akademik" in str(p.name) or "Standar" in t.nama
            results.append({
                "id": t.id,
                "nama": t.nama,
                "deskripsi": t.deskripsi,
                "path": t.path,
                "status": status,
                "missing": missing,
                "not_standalone": not_standalone,
                "unknown": unknown,
                "placeholders": placeholders_list,
                "is_default": is_default,
                "spec": {
                    "paper": "A4",
                    "margin": "4-4-3-3 cm" if is_default else "Sesuai File",
                    "font": "Times New Roman 12pt" if is_default else "Sesuai File",
                    "spacing": "1.5" if is_default else "Sesuai File",
                },
            })
        return results

    def upload_template(self) -> dict[str, Any]:
        if not self._window:
            return {"cancelled": True}
        file_types = ("Word Document (*.docx)", "All files (*.*)")
        res = self._window.create_file_dialog(webview.FileDialog.OPEN, allow_multiple=False, file_types=file_types)
        if not res or len(res) == 0:
            return {"cancelled": True}

        src_path = Path(res[0])
        paths.TEMPLATES_DIR.mkdir(parents=True, exist_ok=True)
        dest_path = paths.TEMPLATES_DIR / src_path.name
        counter = 1
        while dest_path.exists():
            dest_path = paths.TEMPLATES_DIR / f"{src_path.stem}_{counter}{src_path.suffix}"
            counter += 1

        shutil.copy2(src_path, dest_path)
        try:
            val_result = template.validate_template(dest_path)
        except template.TemplateError as exc:
            try:
                dest_path.unlink(missing_ok=True)
            except OSError:
                pass
            return {"error": f"Berkas template tidak valid atau rusak: {exc}"}
        tpl_id = self.db.add_template(database.Template(0, src_path.stem, "Diimpor pengguna", str(dest_path)))

        return {
            "success": True,
            "template": {
                "id": tpl_id,
                "nama": src_path.stem,
                "path": str(dest_path),
                "is_processable": val_result.is_processable,
                "missing": list(val_result.missing),
                "not_standalone": list(val_result.not_standalone),
                "unknown": list(val_result.unknown),
            }
        }

    def delete_template(self, template_id: int) -> dict[str, Any]:
        tpls = self.db.list_template()
        found = next((t for t in tpls if t.id == int(template_id)), None)
        if found and Path(found.path).is_file():
            try:
                Path(found.path).unlink(missing_ok=True)
            except Exception:
                pass
        self.db.delete_template(int(template_id))
        return {"success": True}

    # =========================================================================
    # SOURCES (PDF, WEB, MANUAL)
    # =========================================================================

    def _ingest_single_pdf(self, pdf_path: Path) -> tuple[dict[str, Any] | None, str | None]:
        """Ekstrak teks dan metadata dari satu berkas PDF secara terisolasi.

        Mengembalikan (record, None) jika berhasil, atau (None, error_msg) jika gagal.
        """
        try:
            # Single-pass reading untuk performa optimal
            text, doc_meta_raw = pdf_reader.extract_text_and_metadata(pdf_path)
            doc_meta = normalize_metadata(doc_meta_raw)
        except Exception as e:
            return None, f"{pdf_path.name}: Gagal membaca PDF ({e})"

        try:
            # Unggah cepat: metadata dokumen/DOI cepat; AI di latar
            resolved = intake.fast_doi_source(doc_meta, text[:3000])
            if resolved is not None:
                merged = resolved
            else:
                merged = intake.merge_metadata(Source([], "", ""), doc_meta)
            merged = intake.ensure_title(merged, pdf_path.stem)

            record = intake.make_record(merged, text, "pdf")
            if intake.needs_enrichment(merged):
                self._start_metadata_enrichment(record["ref"], text)
            return record, None
        except Exception as e:
            return None, f"{pdf_path.name}: Gagal memproses metadata ({e})"

    def pick_pdf_source(self) -> dict[str, Any]:
        if not self._window:
            return {"cancelled": True}
        file_types = ("PDF Files (*.pdf)", "All files (*.*)")
        res = self._window.create_file_dialog(webview.FileDialog.OPEN, allow_multiple=True, file_types=file_types)
        if not res or len(res) == 0:
            return {"cancelled": True}

        # Guardrail performa: Maksimal 15 berkas diproses sekaligus
        max_allowed = 15
        selected_paths = [Path(p) for p in res]
        warning_msg: str | None = None
        if len(selected_paths) > max_allowed:
            warning_msg = f"Anda memilih {len(selected_paths)} berkas. Maksimal {max_allowed} berkas diproses sekaligus agar performa tetap ringan."
            selected_paths = selected_paths[:max_allowed]

        records: list[dict[str, Any]] = []
        errors: list[str] = []

        if len(selected_paths) == 1:
            rec, err = self._ingest_single_pdf(selected_paths[0])
            if rec:
                records.append(rec)
            if err:
                errors.append(err)
        else:
            # Ekstraksi paralel multi-threading (maksimal 4 worker threads)
            max_workers = min(len(selected_paths), 4)
            with concurrent.futures.ThreadPoolExecutor(max_workers=max_workers) as executor:
                results = list(executor.map(self._ingest_single_pdf, selected_paths))
                for rec, err in results:
                    if rec:
                        records.append(rec)
                    if err:
                        errors.append(err)

        if not records and errors:
            return {"success": False, "error": "; ".join(errors), "errors": errors}

        return {
            "success": len(records) > 0,
            "record": records[0] if records else None,
            "records": records,
            "count": len(records),
            "errors": errors if errors else None,
            "warning": warning_msg,
        }

    def _start_metadata_enrichment(self, ref: str, text: str) -> None:
        """Ekstraksi bibliografi AI di thread latar; hasil via event.

        Gagal = diam-diam dicatat di log; metadata manual dari pengguna tetap
        menjadi sumber akhir (form koreksi metadata selalu menang).
        """

        def run() -> None:
            with self._enrichment_semaphore:
                try:
                    cl = AIClient(self._cfg)
                    source = bibliography.extract_source(text, cl)
                except (AIError, bibliography.BibliographyError) as exc:
                    log.info(f"Enrichment metadata sumber {ref} gagal: {exc}")
                    return
                except Exception:
                    log.exception(f"Enrichment metadata sumber {ref} error tak terduga")
                    return
                if not (source and (source.authors or source.title)):
                    return
                self._emit("source_metadata_ready", {
                "ref": ref,
                "title": source.title,
                "year": source.year,
                "container": source.container,
                "doi": source.doi,
                "url": source.url,
                "authors": [
                    {
                        "family": a.family,
                        "given": a.given,
                        "is_org": bool(a.is_org),
                        "abbrev": a.abbrev,
                    }
                    for a in source.authors
                ],
            })

        threading.Thread(target=run, daemon=True, name=f"meta-{ref}").start()

    def add_web_source(self, url: str) -> dict[str, Any]:
        url = url.strip()
        if not url:
            return {"success": False, "error": "URL tidak boleh kosong."}
        try:
            text, meta = web_reader.fetch(url)
            meta = normalize_metadata(meta)
        except Exception as e:
            return {"success": False, "error": f"Gagal mengambil halaman web: {e}"}

        # Metadata halaman dipakai seketika; bibliografi AI diperkaya di latar
        # (event source_metadata_ready) supaya menambah sumber terasa instan.
        merged = intake.ensure_title(
            intake.merge_metadata(Source([], "", ""), meta), url
        )
        record = intake.make_record(merged, text, "web")
        if intake.needs_enrichment(merged):
            self._start_metadata_enrichment(record["ref"], text)
        return {"success": True, "record": record}

    def add_manual_source(self, text: str) -> dict[str, Any]:
        text = text.strip()
        if not text:
            return {"error": "Teks sumber tidak boleh kosong."}

        # Jalur cepat DOI bila teks manual memuat DOI; sisanya diperkaya AI
        # di latar. Kebijakan identik dengan PDF/web via sources.intake.
        resolved = intake.fast_doi_source({}, text[:2000])
        if resolved is not None:
            merged = resolved
        else:
            merged = Source([], "", "Teks Manual")

        record = intake.make_record(merged, text, "manual")
        if intake.needs_enrichment(merged):
            self._start_metadata_enrichment(record["ref"], text)
        return {"success": True, "record": record}

    def generate_title_from_sources(self, sources_data: list[dict[str, Any]]) -> dict[str, Any]:
        """Menyusun rekomendasi judul laporan akademik otomatis berdasarkan sumber-sumber jurnal."""
        if not sources_data:
            return {"error": "Tambahkan setidaknya satu sumber jurnal terlebih dahulu."}

        source_summaries = []
        for i, s in enumerate(sources_data, 1):
            title = s.get("title") or s.get("file_name") or f"Sumber {i}"
            authors = s.get("authors") or []
            auth_str = ", ".join(a.get("family", "") for a in authors if a.get("family")) or "Peneliti"
            year = s.get("year", "")
            container = s.get("container", "")
            raw_text = s.get("text", "")
            snippet = " ".join(raw_text[:400].split())
            source_summaries.append(
                f"[{i}] Judul Jurnal: {title}\n"
                f"    Penulis: {auth_str} ({year})\n"
                f"    Jurnal: {container}\n"
                f"    Fokus/Abstrak: {snippet}..."
            )

        sources_text = "\n\n".join(source_summaries)

        system_msg = (
            "Anda adalah dosen pengampu mata kuliah senior dan pakar penulisan karya ilmiah Indonesia. "
            "Tugas Anda adalah merumuskan 1 (SATU) judul laporan akademik, skripsi, makalah penelitian, "
            "atau tugas akhir yang komprehensif, menarik, formal, dan baku dalam Bahasa Indonesia "
            "berdasarkan tema dan sintesis dari sumber-sumber jurnal yang diberikan.\n\n"
            "Aturan pembuatan judul:\n"
            "1. Judul harus mencerminkan sintesis atau fokus penelitian dari jurnal-jurnal tersebut.\n"
            "2. Gunakan gaya bahasa akademik baku (misal: 'Analisis...', 'Implementasi...', 'Kajian Komparatif...', 'Pengaruh...', 'Rancang Bangun...', 'Studi Evaluasi...').\n"
            "3. Panjang judul antara 8 hingga 18 kata.\n"
            "4. Jangan gunakan tanda petik atau awalan seperti 'Judul:' atau 'Berikut adalah judul:'.\n"
            "5. Kembalikan HANYA teks judul laporan tersebut tanpa penjelasan tambahan."
        )

        user_msg = (
            f"Berikut adalah daftar sumber jurnal referensi yang telah ditambahkan:\n\n"
            f"{sources_text}\n\n"
            f"Rumuskan 1 judul laporan akademik terbaik berdasarkan sumber-sumber di atas:"
        )

        cl = AIClient(self._cfg)
        try:
            reply = cl.chat([
                {"role": "system", "content": system_msg},
                {"role": "user", "content": user_msg},
            ])
            candidate = reply.strip().strip("\"'`").replace("\n", " ")
            for prefix in ("Judul:", "Judul Laporan:", "Judul Penelitian:", "Judul Skripsi:"):
                if candidate.lower().startswith(prefix.lower()):
                    candidate = candidate[len(prefix):].strip()

            if candidate:
                return {"success": True, "title": candidate}
        except Exception as exc:
            log.warning(f"Gagal menghasilkan judul via AI: {exc}")

        # Fallback heuristik cerdas bila AI tidak aktif / offline
        first_title = sources_data[0].get("title", "").strip()
        if not first_title or first_title.lower() == "tanpa judul":
            first_title = "Topik Terpilih Berdasarkan Studi Referensi Terkini"

        if len(sources_data) == 1:
            fallback_title = f"Kajian Komprehensif dan Tinjauan Literatur: {first_title}"
        else:
            second_title = sources_data[1].get("title", "").strip()
            if second_title:
                fallback_title = f"Studi Komparatif dan Analisis: {first_title} dan {second_title}"
            else:
                fallback_title = f"Analisis Multiperspektif dan Sintesis Literatur: {first_title}"

        return {"success": True, "title": fallback_title}

    def resolve_doi_source(self, doi: str) -> dict[str, Any]:
        doi = doi.strip()
        if not doi:
            return {"success": False, "error": "DOI tidak boleh kosong."}
        try:
            record = doi_resolver.fetch_doi_metadata(doi)
            return {"success": True, "record": record}
        except Exception as e:
            return {"success": False, "error": f"Gagal mengambil metadata DOI: {e}"}

    # =========================================================================
    # FORMATS & OUTLINES
    # =========================================================================

    def get_format_variants(self) -> list[dict[str, Any]]:
        choices = format_spec.variant_choices()
        res = []
        for code, label in choices:
            var = format_spec.get_variant(code)
            if var:
                res.append({
                    "code": var.code,
                    "name": var.name,
                    "description": var.description,
                    "skeleton": format_spec.skeleton_lines(var),
                    "outline_titles": format_spec.outline_titles(var),
                })
            else:
                res.append({
                    "code": "bebas",
                    "name": "Bebas",
                    "description": "Outline manual bebas tanpa struktur baku.",
                    "skeleton": ["Bebas menyusun outline."],
                    "outline_titles": [],
                })
        return res

    def suggest_outline(self, topic: str, sources_data: list[dict[str, Any]]) -> dict[str, Any]:
        topic = topic.strip()
        if not topic:
            return {"error": "Tulis judul laporan terlebih dahulu."}

        sources = [record_to_source(s) for s in sources_data]
        texts = [s.get("text", "") for s in sources_data]
        cl = AIClient(self._cfg)
        try:
            outlines = content.suggest_outline(topic, sources, texts, cl)
            return {"success": True, "outline": outlines}
        except Exception as e:
            return {"error": f"Gagal menyusun outline: {e}"}

    # =========================================================================
    # DRAFTS
    # =========================================================================

    def load_active_draft(self) -> dict[str, Any]:
        draft = drafts.load_draft(DRAFT_ACTIVE_ID)
        if not draft:
            return {"has_draft": False}
        fields = dict(draft.fields) if draft.fields else {}
        profil = self.db.get_profil()
        if profil:
            if (profil.nama or "").strip():
                fields["NAMA"] = profil.nama.strip()
            if (profil.nim or "").strip():
                fields["NIM"] = profil.nim.strip()
            if (profil.prodi or "").strip():
                fields["PRODI"] = profil.prodi.strip()
            if (profil.fakultas or "").strip():
                fields["FAKULTAS"] = profil.fakultas.strip()
            if (profil.institusi or "").strip():
                fields["INSTITUSI"] = profil.institusi.strip()
        return {
            "has_draft": True,
            "draft_id": draft.draft_id,
            "fields": fields,
            "sources": draft.sources,
            "sections": draft.sections,
            "done": list(draft.done),
        }

    def save_active_draft(self, state: dict[str, Any]) -> dict[str, Any]:
        # Kebijakan merge/preserve/clear draf aktif dimiliki drafts.py
        # (lihat drafts.reconcile_draft) dan diuji langsung di sana; adapter
        # ini hanya menerjemahkan panggilan pywebview.
        drafts.reconcile_draft(DRAFT_ACTIVE_ID, state)
        return {"success": True}

    def clear_active_draft(self) -> dict[str, Any]:
        # Arsipkan draf aktif yang masih berisi sebelum dihapus, agar aliran
        # "Mulai Lembar Kerja Kosong Baru" tidak memusnahkan pekerjaan lama.
        existing = drafts.load_draft(DRAFT_ACTIVE_ID)
        if existing and drafts.has_content(existing):
            drafts.archive_draft(existing, prefix="arsip")
        drafts.delete_draft(DRAFT_ACTIVE_ID)
        return {"success": True}

    def list_saved_drafts(self) -> list[dict[str, Any]]:
        draft_ids = drafts.list_drafts()
        results = []
        for d_id in draft_ids:
            st = drafts.load_draft(d_id)
            if not st:
                continue
            title = st.fields.get("JUDUL", "Tanpa Judul") if st.fields else "Tanpa Judul"
            sentence_count = sum(len(p.get("sentences", [])) for sec in st.sections for p in sec.get("paragraphs", []))
            try:
                p = paths.DRAFTS_DIR / f"{d_id}.json"
                mtime = p.stat().st_mtime
            except Exception:
                mtime = 0
            results.append({
                "draft_id": d_id,
                "title": title or "Tanpa Judul",
                "sources_count": len(st.sources),
                "sections_count": len(st.sections),
                "sentence_count": sentence_count,
                "is_active": d_id == DRAFT_ACTIVE_ID,
                "mtime": mtime,
            })
        results.sort(key=lambda x: (not x["is_active"], -x["mtime"]))
        return results

    def save_named_draft(self, name: str, state: dict[str, Any]) -> dict[str, Any]:
        try:
            if not isinstance(state, dict):
                return {"success": False, "error": "Data state draf harus berupa objek."}
            slug = report.sanitize_filename(name, max_len=40)
            if not slug or slug == "laporan":
                slug = f"laporan_{report.sanitize_filename(state.get('fields', {}).get('JUDUL', 'proyek'), max_len=20)}"
            draft_state = drafts.DraftState(
                draft_id=slug,
                fields=state.get("fields", {}),
                sources=state.get("sources", []),
                sections=state.get("sections", []),
                done=list(state.get("done", [])),
            )
            drafts.save_draft(draft_state)
            return {"success": True, "draft_id": slug}
        except Exception as e:
            return {"success": False, "error": str(e)}

    def load_named_draft(self, draft_id: str) -> dict[str, Any]:
        st = drafts.load_draft(draft_id)
        if not st:
            return {"error": "Draf tidak ditemukan."}
        active_copy = drafts.DraftState(
            draft_id=DRAFT_ACTIVE_ID,
            fields=st.fields,
            sources=st.sources,
            sections=st.sections,
            done=st.done,
        )
        drafts.save_draft(active_copy)
        return {
            "success": True,
            "draft": {
                "has_draft": True,
                "draft_id": st.draft_id,
                "fields": st.fields,
                "sources": st.sources,
                "sections": st.sections,
                "done": list(st.done),
            }
        }

    def delete_named_draft(self, draft_id: str) -> dict[str, Any]:
        drafts.delete_draft(draft_id)
        return {"success": True}

    def delete_named_drafts(self, draft_ids: list[str]) -> dict[str, Any]:
        for d_id in draft_ids:
            if d_id:
                drafts.delete_draft(str(d_id))
        return {"success": True}

    # =========================================================================
    # GENERATION ORCHESTRATION (STREAMING VIA JS EVENTS)
    # =========================================================================

    def start_generation(self, request_data: dict[str, Any]) -> dict[str, Any]:
        if self._gen_thread and self._gen_thread.is_alive():
            return {"error": "Generasi sedang berjalan."}

        topic = request_data.get("topic", "").strip()
        outline = request_data.get("outline", [])
        sources_data = request_data.get("sources", [])
        variant_code = request_data.get("variant", "d")

        if not topic:
            return {"error": "Judul belum diisi."}
        if not sources_data:
            return {"error": "Masukkan minimal satu sumber referensi."}

        variant = format_spec.get_variant(variant_code)
        if not outline:
            if variant:
                outline = format_spec.outline_titles(variant)
            else:
                outline = [
                    "BAB 1. PENDAHULUAN",
                    "BAB 2. TINJAUAN PUSTAKA",
                    "BAB 3. HASIL DAN PEMBAHASAN",
                    "BAB 4. KESIMPULAN DAN SARAN",
                ]

        self._cancel_gen_event.clear()

        sources = [record_to_source(s) for s in sources_data]
        texts = [s.get("text", "") for s in sources_data]
        variant = format_spec.get_variant(variant_code)

        class Progress(content.GenerationProgress):
            def __init__(self, bridge: ApiBridge):
                self.bridge = bridge

            def batch_started(self, n: int, total: int) -> None:
                self.bridge._emit("gen_batch_progress", {"batch": n, "total": total})

            def sentence_born(self, sentence: content.Sentence, status: str) -> None:
                self.bridge._emit("gen_sentence_born", {
                    "uid": sentence.uid,
                    "text": sentence.text,
                    "status": status,
                    "is_factual": sentence.is_factual,
                    "source_ids": sentence.source_ids,
                })

            def finished(self, cancelled: bool) -> None:
                self.bridge._emit("gen_finished", {"cancelled": cancelled})

        def run_thread():
            prog = Progress(self)
            cl = AIClient(self._cfg)
            try:
                sections = content.generate_sections(
                    topic=topic,
                    outline=outline,
                    sources=sources,
                    source_texts=texts,
                    client=cl,
                    progress=prog,
                    cancel=self._cancel_gen_event,
                    variant=variant,
                )
                serialized = report.sections_to_dicts(sections)
                statuses = verify_sections(sections, texts)
                serialized_statuses = [
                    {
                        "index": st.index,
                        "uid": st.sentence.uid,
                        "status": st.status,
                        "reason": getattr(st, "reason", ""),
                        "is_factual": st.sentence.is_factual,
                        "source_ids": st.sentence.source_ids,
                    }
                    for st in statuses
                ]
                active_draft = drafts.load_draft(DRAFT_ACTIVE_ID)
                draft_fields = dict(active_draft.fields) if active_draft and active_draft.fields else {}
                draft_fields["JUDUL"] = topic
                draft_fields["VARIANT"] = variant_code
                draft_fields["OUTLINE"] = outline
                profil = self.db.get_profil()
                if profil:
                    if (profil.nama or "").strip():
                        draft_fields["NAMA"] = profil.nama.strip()
                    if (profil.nim or "").strip():
                        draft_fields["NIM"] = profil.nim.strip()
                    if (profil.prodi or "").strip():
                        draft_fields["PRODI"] = profil.prodi.strip()
                    if (profil.fakultas or "").strip():
                        draft_fields["FAKULTAS"] = profil.fakultas.strip()
                    if (profil.institusi or "").strip():
                        draft_fields["INSTITUSI"] = profil.institusi.strip()
                if active_draft and drafts.has_review_progress(active_draft):
                    drafts.archive_draft(active_draft, prefix="arsip_pregen")
                drafts.save_draft(drafts.DraftState(
                    draft_id=DRAFT_ACTIVE_ID,
                    fields=draft_fields,
                    sources=sources_data,
                    sections=serialized,
                    done=[],
                ))
                self._emit("gen_complete", {
                    "cancelled": False,
                    "sections": serialized,
                    "statuses": serialized_statuses,
                })
            except Exception as e:
                log.exception("Generation failed")
                self._emit("gen_error", {"error": str(e)})

        self._gen_thread = threading.Thread(target=run_thread, daemon=True)
        self._gen_thread.start()
        return {"started": True}

    def cancel_generation(self) -> dict[str, Any]:
        self._cancel_gen_event.set()
        return {"success": True}

    # =========================================================================
    # REVIEW & EDITING
    # =========================================================================

    def verify_document(self, sections_dicts: list[dict], sources_data: list[dict]) -> list[dict]:
        if not sources_data:
            # Tanpa sumber, semua kalimat faktual otomatis gagal verifikasi.
            # Jatuhkan ke sumber tersimpan pada draf aktif agar status tetap
            # konsisten dengan yang terlihat saat streaming.
            draft = drafts.load_draft(DRAFT_ACTIVE_ID)
            if draft and draft.sources:
                sources_data = draft.sources
        return reviewed.ReviewedDocument.from_payload(sections_dicts, sources_data).status_payload()

    def mark_non_factual_done(self, sections_dicts: list[dict], current_done: list[str]) -> dict[str, Any]:
        sections = report.sections_from_dicts(sections_dicts)
        done_set = set(current_done)
        for sec in sections:
            for p in sec.paragraphs:
                for s in p.sentences:
                    if not s.is_factual:
                        done_set.add(s.uid)
        return {"success": True, "done": list(done_set)}

    def paraphrase_sentence(self, sentence_dict: dict, sources_data: list[dict]) -> dict[str, Any]:
        sent = report.sentence_from_dict(sentence_dict)
        para = content.Paragraph([sent])
        sources = [record_to_source(s) for s in sources_data]
        texts = [s.get("text", "") for s in sources_data]
        cl = AIClient(self._cfg)
        try:
            new_para = content.paraphrase(para, sources, texts, cl)
            if new_para.sentences:
                new_sent = new_para.sentences[0]
                return {
                    "success": True,
                    "sentence": report.sentence_to_dict(new_sent),
                }
            return {"error": "Hasil parafrase kosong."}
        except Exception as e:
            return {"error": f"Gagal parafrase: {e}"}

    def request_dosen_review(self, sections_dicts: list[dict], sources_data: list[dict]) -> dict[str, Any]:
        sections = report.sections_from_dicts(sections_dicts)
        sources = [record_to_source(s) for s in sources_data]
        cited = {i for sec in sections for p in sec.paragraphs for s in p.sentences for i in s.source_ids}
        refs = report.references_for(sources, cited)
        cl = AIClient(self._cfg)
        try:
            notes = dosen.review_report(sections, refs, cl)
            return {
                "success": True,
                "notes": [asdict(n) for n in notes],
            }
        except Exception as e:
            return {"error": f"Gagal review dosen: {e}"}

    # =========================================================================
    # CITATION LABELS & BIBLIOGRAPHY
    # =========================================================================

    def compute_citations_and_references(
        self, sources_data: list[dict], sections_dicts: list[dict], style: str = "et_al"
    ) -> dict[str, Any]:
        return reviewed.ReviewedDocument.from_payload(sections_dicts, sources_data).citation_payload(style)

    def get_reviewed_snapshot(
        self, sections_dicts: list[dict], sources_data: list[dict], style: str | None = None
    ) -> dict[str, Any]:
        """Satu panggilan untuk layar review & sidebar: status kalimat, label
        sitasi, daftar pustaka, dan analitik — tanpa tiga round-trip terpisah."""
        doc = reviewed.ReviewedDocument.from_payload(sections_dicts, sources_data)
        return doc.snapshot(style or self._cfg.citation_style)

    def get_document_analytics(
        self, sections_dicts: list[dict], sources_data: list[dict]
    ) -> dict[str, Any]:
        """Hitung metrik kualitas dokumen dan analitik mendalam standar akademik SaaS."""
        return reviewed.ReviewedDocument.from_payload(sections_dicts, sources_data).analytics()

    def open_output_folder(self) -> dict[str, Any]:
        """Buka folder keluaran dokumen DOCX di Windows Explorer."""
        try:
            paths.OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
            if os.name == "nt":
                os.startfile(str(paths.OUTPUT_DIR.resolve()))
            return {"success": True}
        except Exception as e:
            return {"success": False, "error": f"Gagal membuka folder output: {e}"}

    # =========================================================================
    # PREFLIGHT & EXPORT
    # =========================================================================

    def run_preflight(self, state: dict[str, Any]) -> dict[str, Any]:
        sources_data = state.get("sources", [])
        sections_dicts = state.get("sections", [])
        done = set(state.get("done", []))
        fields = state.get("fields", {})

        doc = reviewed.ReviewedDocument.from_payload(sections_dicts, sources_data)
        judul = fields.get("JUDUL", "Laporan")
        logo_path = Path(fields.get("LOGO_PATH")) if fields.get("LOGO_PATH") else None
        cited = doc.cited_ids
        style = self._cfg.citation_style
        done_set = frozenset(done)

        checks = preflight.run_checks(
            judul=judul,
            sources=doc.sources,
            statuses=doc.statuses,
            cited=cited,
            logo_path=logo_path,
            style=style,
            done=done_set,
            sections=doc.sections,
        )
        issues = preflight.find_issues(
            judul=judul,
            sources=doc.sources,
            statuses=doc.statuses,
            cited=cited,
            logo_path=logo_path,
            style=style,
            done=done_set,
            sections=doc.sections,
        )

        return {
            "checks": [asdict(c) for c in checks],
            "issues": [asdict(i) for i in issues],
            "can_export_directly": len(issues) == 0,
        }

    def export_report(self, state: dict[str, Any]) -> dict[str, Any]:
        sources_data = state.get("sources", [])
        sections_dicts = state.get("sections", [])
        fields = state.get("fields", {})
        template_id = state.get("template_id")
        open_word = state.get("open_word", True)
        style = state.get("style", self._cfg.citation_style)

        tpl_id_int: int | None = None
        if template_id is not None:
            try:
                tpl_id_int = int(template_id)
            except (ValueError, TypeError):
                tpl_id_int = None
        if tpl_id_int is None:
            return {"error": "Pilih template terlebih dahulu"}

        tpls = self.db.list_template()
        tpl = next((t for t in tpls if t.id == tpl_id_int), None)
        # Fallback toleran jika id template out-of-sync tetapi ada template valid di database
        if not tpl and tpls:
            tpl = next((t for t in tpls if Path(t.path).is_file()), None)
        if not tpl or not Path(tpl.path).is_file():
            return {"error": "Template DOCX tidak ditemukan."}

        doc = reviewed.ReviewedDocument.from_payload(sections_dicts, sources_data)
        cited = doc.cited_ids

        labels = report.labels_for(doc.sources, cited, style=style)
        label_forms = report.label_forms_for(doc.sources, cited, style=style)
        refs = report.references_for(doc.sources, cited)

        # Gerbang ekspor keras: dokumen yang belum layak dikumpulkan ditolak di
        # backend, bukan hanya lewat dialog preflight yang bisa dilewati. Biaya
        # O(jumlah kalimat) tanpa verifikasi ulang dan tanpa panggilan AI.
        blockers = preflight.export_blockers(doc.sections, refs)
        if blockers:
            return {
                "error": "Ekspor ditolak — dokumen belum layak dikumpulkan: "
                + "; ".join(blockers)
                + ". Lengkapi dokumen (atau buang penanda yang tidak dipakai) lalu ekspor lagi."
            }

        normalized_fields = dict(fields)
        for k, v in list(fields.items()):
            if not k.startswith("{{") and not k.endswith("}}"):
                normalized_fields.setdefault(f"{{{{{k}}}}}", v)
            elif k.startswith("{{") and k.endswith("}}"):
                normalized_fields.setdefault(k[2:-2].strip(), v)

        # Tanggal dari PDF/unduhan/skrip uji bisa berbahasa Inggris; normalkan ke
        # bentuk Indonesia sebelum dirender ke sampul.
        for key in list(normalized_fields):
            if key.strip("{}").strip().upper() == "TANGGAL":
                normalized_fields[key] = report.format_tanggal_teks(normalized_fields[key])

        # Profil mahasiswa di Tab Data (database SQLite) adalah SUMBER UTAMA
        # (single source of truth) untuk halaman sampul. Bila profil di database
        # terisi, nilai tersebut SELALU digunakan agar residu data draf lama/uji
        # tidak pernah mengesampingkan identitas resmi mahasiswa.
        profil = self.db.get_profil()
        official_identity = {
            "NAMA": (profil.nama or "").strip(),
            "NIM": (profil.nim or "").strip(),
            "PRODI": (profil.prodi or "").strip().upper(),
            "FAKULTAS": (profil.fakultas or "").strip().upper(),
            "INSTITUSI": (profil.institusi or "").strip().upper(),
        }
        for key, value in official_identity.items():
            if value:
                normalized_fields[key] = value
                normalized_fields[f"{{{{{key}}}}}"] = value

        # Tanggal dan tahun otomatis bila belum dispesifikasi
        if not str(normalized_fields.get("TANGGAL", "") or "").strip():
            tgl = report.format_tanggal()
            normalized_fields["TANGGAL"] = tgl
            normalized_fields["{{TANGGAL}}"] = tgl
        if not str(normalized_fields.get("TAHUN", "") or "").strip():
            yr = str(datetime.now().year)
            normalized_fields["TAHUN"] = yr
            normalized_fields["{{TAHUN}}"] = yr

        # Pastikan PRODI, FAKULTAS, INSTITUSI selalu kapital rapi untuk sampul resmi
        for key in ("PRODI", "FAKULTAS", "INSTITUSI"):
            for field_key in (key, f"{{{{{key}}}}}"):
                val = str(normalized_fields.get(field_key, "") or "").strip()
                if val:
                    normalized_fields[field_key] = val.upper()

        judul = normalized_fields.get("JUDUL", "Laporan")
        filename = report.sanitize_filename(judul)
        paths.OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
        out_path = paths.OUTPUT_DIR / f"{filename}.docx"

        logo_path = normalized_fields.get("LOGO_PATH")
        if logo_path and not Path(logo_path).is_file():
            logo_path = None

        try:
            # Daftar pustaka dikirim terpejal per run (teks, italic) agar renderer
            # bisa memiringkan nama jurnal + volume APA 7; fallback string bila kosong.
            ref_runs = report.reference_runs_for(doc.sources, cited) or refs
            final_path = renderer.render_report(
                template_path=Path(tpl.path),
                fields=normalized_fields,
                sections=doc.sections,
                references=ref_runs,
                logo_path=Path(logo_path) if logo_path else None,
                output_path=out_path,
                labels=labels,
                update_fields=True,
                label_forms=label_forms,
            )
            leftover = _leftover_placeholders(final_path)
            if leftover:
                try:
                    final_path.unlink()
                except OSError:
                    pass
                return {
                    "error": (
                        "Ekspor dibatalkan: penanda yang tidak terganti masih ada "
                        f"di dokumen ({', '.join(leftover)}). Lengkapi data yang "
                        "bersangkutan lalu ekspor lagi."
                    )
                }
            # Simpan draf (jangan dihapus agar dapat dibuka & direvisi kembali)
            # drafts.delete_draft(DRAFT_ACTIVE_ID)

            if open_word and os.name == "nt":
                try:
                    os.startfile(str(final_path))
                except Exception:
                    pass

            return {"success": True, "output_path": str(final_path)}
        except Exception as e:
            log.exception("Export failed")
            return {"error": f"Gagal mengekspor laporan: {e}"}
