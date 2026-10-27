"""Pencarian dan penggantian penanda `{{...}}` lintas run (9.1, 9.2, 9.4)."""

from __future__ import annotations

import re
from typing import Iterator

from docx.oxml.ns import qn
from docx.text.run import Run

PLACEHOLDER_RE = re.compile(r"\{\{[^{}]+\}\}")
_XML_ILLEGAL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")


def clean_text(text: str) -> str:
    """Buang karakter kontrol yang ditolak XML dan karakter glyph rusak."""
    text = re.sub(r"(\w+s)[\ufffd](\s|$)", r"\1'\2", text, flags=re.I)
    text = re.sub(r"(\w+)[\ufffd](s|d|t|re|ve|ll|m)\b", r"\1'\2", text, flags=re.I)
    text = text.replace("\ufffd", "").replace("", "")
    return _XML_ILLEGAL.sub("", text)


def find_names(text: str) -> list[str]:
    """Semua penanda yang muncul di `text`, berurutan, termasuk yang berulang."""
    return [match.group(0) for match in PLACEHOLDER_RE.finditer(text)]


def _paragraphs_in(container) -> Iterator:
    for paragraph in container.paragraphs:
        yield paragraph
    for table in container.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    yield paragraph


def iter_paragraphs(document) -> Iterator:
    """Semua paragraf: badan dokumen, sel tabel, header, dan footer (9.2).

    Header/footer yang berbeda untuk halaman pertama dan halaman genap ikut
    dipindai; tanpa itu penanda di dalamnya lolos validasi lalu tercetak mentah.
    Text box dan shape tidak didukung dan dicatat di README (Task 18).
    """
    yield from _paragraphs_in(document)
    for section in document.sections:
        for container in (
            section.header,
            section.first_page_header,
            section.even_page_header,
            section.footer,
            section.first_page_footer,
            section.even_page_footer,
        ):
            yield from _paragraphs_in(container)


def collect_names(document) -> list[str]:
    """Nama penanda unik di seluruh dokumen, berurutan sesuai kemunculan pertama."""
    names: list[str] = []
    seen: set[str] = set()
    for paragraph in iter_paragraphs(document):
        for name in find_names(paragraph.text):
            if name not in seen:
                seen.add(name)
                names.append(name)
    return names


def _run_bounds(runs) -> list[tuple[int, int]]:
    bounds: list[tuple[int, int]] = []
    position = 0
    for run in runs:
        bounds.append((position, position + len(run.text)))
        position += len(run.text)
    return bounds


def _replace_span(runs, bounds: list[tuple[int, int]], start: int, end: int, replacement: str) -> None:
    """Ganti `text[start:end]` lintas run; format run pertama dipertahankan (9.1).

    Teks sebelum penanda tetap di run pertama, teks sesudahnya tetap di run
    terakhir, run yang seluruhnya di dalam penanda dikosongkan.
    """
    first = None
    last = None
    for index, (run_start, run_end) in enumerate(bounds):
        if first is None and start < run_end and end > run_start:
            first = index
        if run_start < end <= run_end:
            last = index
    if first is None or last is None:
        return
    if last == first:
        first_start = bounds[first][0]
        run = runs[first]
        run.text = run.text[: start - first_start] + replacement + run.text[end - first_start :]
        return
    first_run = runs[first]
    first_run.text = first_run.text[: start - bounds[first][0]] + replacement
    for index in range(first + 1, last):
        runs[index].text = ""
    last_run = runs[last]
    last_run.text = last_run.text[end - bounds[last][0] :]


def replace_in_paragraph(paragraph, mapping: dict[str, str]) -> list[str]:
    """Ganti penanda di `mapping`, termasuk yang terpotong antar run (9.1).

    Mengembalikan nama penanda yang benar-benar tergantikan, berurutan dari kiri
    ke kanan. Penanda yang tidak ada di `mapping` dibiarkan apa adanya.
    """
    runs = all_runs(paragraph)
    if not runs:
        return []
    combined = "".join(run.text for run in runs)
    matches = [match for match in PLACEHOLDER_RE.finditer(combined) if match.group(0) in mapping]
    if not matches:
        return []
    replaced: list[str] = []
    for match in reversed(matches):
        _replace_span(
            runs,
            _run_bounds(runs),
            match.start(),
            match.end(),
            clean_text(str(mapping[match.group(0)])),
        )
        replaced.append(match.group(0))
    replaced.reverse()
    return replaced


def is_standalone(paragraph, name: str) -> bool:
    """True bila paragraf hanya berisi satu penanda tanpa teks lain (9.4)."""
    return paragraph.text.strip() == name


def all_runs(paragraph) -> list[Run]:
    """Semua run di paragraf, termasuk yang berada di dalam hyperlink.

    `paragraph.runs` hanya mengembalikan run anak langsung; penanda di dalam
    hyperlink karenanya terlihat oleh validasi tetapi tidak pernah terganti.
    """
    return [Run(element, paragraph) for element in paragraph._p.iter(qn("w:r"))]
