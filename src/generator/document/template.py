"""Validasi template DOCX sebelum dipakai (9.2, 9.3, 9.4, 12.5)."""

from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path

from docx import Document

from . import placeholders

# Sumber: desain 9, tabel penanda; {{DAFTAR_ISI}} wajib supaya daftar isi selalu terpenuhi.
REQUIRED_PLACEHOLDERS: list[str] = [
    "{{JUDUL}}",
    "{{NAMA}}",
    "{{NIM}}",
    "{{MATA_KULIAH}}",
    "{{DOSEN}}",
    "{{TANGGAL}}",
    "{{ISI_LAPORAN}}",
    "{{DAFTAR_PUSTAKA}}",
    "{{DAFTAR_ISI}}",
    "{{LOGO}}",
]

# Diganti bukan sekadar teks: blok paragraf, pustaka, field TOC, atau gambar 5x5 cm.
SPECIAL_PLACEHOLDERS: list[str] = [
    "{{ISI_LAPORAN}}",
    "{{DAFTAR_PUSTAKA}}",
    "{{DAFTAR_ISI}}",
    "{{LOGO}}",
]

OPTIONAL_PLACEHOLDERS: list[str] = ["{{TAHUN}}", "{{PRODI}}"]

_KNOWN = set(REQUIRED_PLACEHOLDERS) | set(SPECIAL_PLACEHOLDERS) | set(OPTIONAL_PLACEHOLDERS)


class TemplateError(Exception):
    """Template tidak bisa dibuka atau bukan berkas .docx yang valid."""


@dataclass
class ValidationResult:
    missing: list[str] = field(default_factory=list)
    not_standalone: list[str] = field(default_factory=list)
    unknown: list[str] = field(default_factory=list)

    @property
    def is_processable(self) -> bool:
        return not self.missing and not self.not_standalone


def load_template(path: Path):
    """Buka template; kegagalan apa pun dibungkus TemplateError (12.5)."""
    try:
        return Document(path)
    except Exception as error:
        raise TemplateError("File template tidak valid atau bukan .docx") from error


def validate_template(path: Path) -> ValidationResult:
    """Periksa penanda wajib, penanda khusus yang tidak berdiri sendiri, dan tak dikenal."""
    return validate_document(load_template(path))


def validate_document(document) -> ValidationResult:
    """Validasi isi dokumen yang sudah dibuka (dipakai juga oleh renderer)."""
    present = placeholders.collect_names(document)
    present_set = set(present)

    missing = [name for name in REQUIRED_PLACEHOLDERS if name not in present_set]

    not_standalone: list[str] = []
    for name in SPECIAL_PLACEHOLDERS:
        if name not in present_set:
            continue
        standalone_somewhere = any(
            name in placeholders.find_names(paragraph.text)
            and placeholders.is_standalone(paragraph, name)
            for paragraph in placeholders.iter_paragraphs(document)
        )
        if not standalone_somewhere:
            not_standalone.append(name)

    unknown = [name for name in present if name not in _KNOWN]
    return ValidationResult(missing=missing, not_standalone=not_standalone, unknown=unknown)
