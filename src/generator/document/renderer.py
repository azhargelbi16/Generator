"""Renderer DOCX: isi penanda, logo 5x5 cm, field TOC, heading berlevel, hanging indent (18, 20)."""

from __future__ import annotations

import os
import re
import tempfile
from datetime import datetime
from pathlib import Path

from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Cm, Emu, Pt
from docx.text.paragraph import Paragraph as DocxParagraph

from ..ai import content
from ..logging_setup import get_logger
from . import placeholders, template

_log = get_logger("renderer")

TOC_FIELD_XML = (
    f"<w:p {nsdecls('w')}>"
    '<w:r><w:fldChar w:fldCharType="begin"/></w:r>'
    '<w:r><w:instrText xml:space="preserve"> TOC \\o "1-9" \\h \\z \\u </w:instrText></w:r>'
    '<w:r><w:fldChar w:fldCharType="separate"/></w:r>'
    '<w:r><w:fldChar w:fldCharType="end"/></w:r>'
    "</w:p>"
)

_LOGO_BOX = Cm(5)
_HANGING = Cm(1.27)
_MAX_HEADING_LEVEL = 9


def _heading_style(doc, level: int) -> str:
    """Nama style Heading yang aman dipakai; dibuat bila template tidak memilikinya (18.1)."""
    level = min(max(level, 1), _MAX_HEADING_LEVEL)
    name = f"Heading {level}"
    if name in doc.styles:
        return name
    try:
        style = doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        style.base_style = doc.styles["Normal"]
        style.font.bold = True
        return name
    except Exception:
        return "Normal"


def _insert_paragraph_before(anchor: DocxParagraph, text: str = "", style: str | None = None) -> DocxParagraph:
    new_p = OxmlElement("w:p")
    anchor._p.addprevious(new_p)
    paragraph = DocxParagraph(new_p, anchor._parent)
    if style:
        paragraph.style = style
    if text:
        paragraph.add_run(placeholders.clean_text(text))
    return paragraph


def _remove_paragraph(paragraph: DocxParagraph) -> None:
    element = paragraph._p
    element.getparent().remove(element)


def _clear_runs(paragraph: DocxParagraph) -> None:
    for run in paragraph.runs:
        run.text = ""


def _logo_size(logo_path: Path) -> tuple[Emu, Emu]:
    """Skala agar masuk kotak 5x5 cm tanpa memotong, rasio asli terjaga.

    Orientasi EXIF diterapkan lebih dulu: foto ponsel sering tersimpan miring
    dan Word memakai dimensi setelah rotasi.
    """
    from PIL import Image, ImageOps

    with Image.open(logo_path) as image:
        width, height = ImageOps.exif_transpose(image).size
    if width <= 0 or height <= 0:
        raise ValueError("ukuran gambar tidak valid")
    scale = min(_LOGO_BOX / width, _LOGO_BOX / height)
    return Emu(int(round(width * scale))), Emu(int(round(height * scale)))


def _place_logo(paragraph: DocxParagraph, logo_path: Path | None) -> None:
    # Tanpa logo, paragraf dihapus total: menyisakan paragraf kosong hanya
    # membuat lubang vertikal (margin 2 cm) di tengah halaman sampul.
    if logo_path is None or not Path(logo_path).is_file():
        _remove_paragraph(paragraph)
        return
    _clear_runs(paragraph)
    try:
        width, height = _logo_size(logo_path)
        run = paragraph.runs[0] if paragraph.runs else paragraph.add_run()
        run.add_picture(str(logo_path), width=width, height=height)
    except Exception as error:
        _log.warning("Logo dilewati karena tidak bisa dimuat (%s): %s", logo_path, error)
        _remove_paragraph(paragraph)
        return
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER


def _place_toc(paragraph: DocxParagraph) -> None:
    paragraph._p.addprevious(parse_xml(TOC_FIELD_XML))
    _remove_paragraph(paragraph)


def _enable_update_fields(doc) -> None:
    """Minta Word memperbarui field (daftar isi) saat dokumen pertama dibuka."""
    settings = doc.settings.element
    if settings.find(qn("w:updateFields")) is not None:
        return
    element = OxmlElement("w:updateFields")
    element.set(qn("w:val"), "true")
    # Skema OOXML menuntut updateFields sebelum compat/rsids; Word mengabaikannya bila urutannya salah.
    compat = settings.find(qn("w:compat"))
    if compat is not None:
        compat.addprevious(element)
    else:
        settings.append(element)


def _footer_has_page_field(footer) -> bool:
    for instr in footer._element.iter(qn("w:instrText")):
        if "PAGE" in (instr.text or "").upper():
            return True
    return False


def _footer_has_text(footer) -> bool:
    return any(paragraph.text.strip() for paragraph in footer.paragraphs)


def _set_page_number_start(section, start: int) -> None:
    """Atur nomor halaman pertama section (0 = sampul tidak bernomor)."""
    sect_pr = section._sectPr
    existing = sect_pr.find(qn("w:pgNumType"))
    if existing is not None:
        existing.set(qn("w:start"), str(start))
        return
    element = OxmlElement("w:pgNumType")
    element.set(qn("w:start"), str(start))
    # Urutan elemen sectPr diatur skema OOXML; sisipkan sebelum elemen sesudahnya.
    sect_pr.insert_element_before(
        element,
        "w:cols", "w:formProt", "w:vAlign", "w:noEndnote", "w:titlePg",
        "w:textDirection", "w:bidi", "w:rtlGutter", "w:docGrid",
        "w:printerSettings", "w:sectPrChange",
    )


def _add_page_number_run(paragraph) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    number = OxmlElement("w:t")
    number.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for node in (begin, instr, separate, number, end):
        run._r.append(node)


def _ensure_page_numbers(doc) -> None:
    """Nomor halaman bawah-tengah bila template belum menyediakannya.

    Footer yang sudah berisi teks atau field PAGE dibiarkan: itu milik template
    atau pengguna. Sampul (halaman pertama section) tidak diberi nomor lewat
    titlePg + start 0, sehingga halaman isi pertama tampil sebagai halaman 1.
    """
    for section in doc.sections:
        footer = section.footer
        if _footer_has_page_field(footer) or _footer_has_text(footer):
            continue
        section.different_first_page_header_footer = True
        _set_page_number_start(section, 0)
        paragraph = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _add_page_number_run(paragraph)


# Pola awal butir daftar: "1. ", "2) ", "a. " (maks dua digit agar kalimat
# naratif yang kebetulan diawali tahun tidak ikut terpecah).
_NUMBERED_ITEM_RE = re.compile(r"^\s*(?:\d{1,2}|[a-z])[\.\)]\s+\S")


def _numbered_groups(block) -> list[list] | None:
    """Kelompokkan kalimat blok menjadi paragraf: satu butir satu paragraf.

    Kembalikan ``None`` bila blok bukan daftar (kurang dari dua butir bernomor),
    sehingga paragraf naratif biasa tetap dirender utuh. Kalimat pengantar yang
    bukan butir tetap dikelompokkan bersama; setiap butir berdiri sendiri.
    """
    groups: list[list] = []
    current: list = []
    items = 0
    for sentence in block.sentences:
        text = sentence.text.strip()
        if not text:
            continue
        if _NUMBERED_ITEM_RE.match(text):
            if current:
                groups.append(current)
                current = []
            groups.append([sentence])
            items += 1
        else:
            current.append(sentence)
    if current:
        groups.append(current)
    return groups if items >= 2 else None


def _place_body(doc, paragraph: DocxParagraph, sections, labels, label_forms=None) -> None:
    defined: set[int] = set()
    for i, section in enumerate(sections):
        heading_style = _heading_style(doc, section.level)
        heading = _insert_paragraph_before(paragraph, section.title, heading_style)
        if heading_style == "Normal":
            for run in heading.runs:
                run.bold = True

        # Standar akademik: setiap Bab (level 1) dimulai di halaman baru.
        # Bab pertama tidak perlu karena template sudah memberi halaman
        # pemisah sebelum {{ISI_LAPORAN}}.
        if section.level == 1:
            if i > 0:
                heading.paragraph_format.page_break_before = True
            heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
            heading.paragraph_format.space_before = Pt(12)
            heading.paragraph_format.space_after = Pt(12)
        else:
            heading.alignment = WD_ALIGN_PARAGRAPH.LEFT
            heading.paragraph_format.space_before = Pt(12)
            heading.paragraph_format.space_after = Pt(6)

        # Cegah duplikasi teks jika bab level 1 memiliki paragraf yang sama persis
        # dengan subbab level 2 tepat di bawahnya (misal BAB 1 dan 1.1 Latar Belakang).
        paragraphs = section.paragraphs
        if section.level == 1 and i + 1 < len(sections) and sections[i + 1].level == 2:
            next_p = sections[i + 1].paragraphs
            if paragraphs and next_p:
                cur_texts = [s.text.strip() for p in paragraphs for s in p.sentences if s.text.strip()]
                next_texts = [s.text.strip() for p in next_p for s in p.sentences if s.text.strip()]
                if cur_texts and cur_texts == next_texts:
                    paragraphs = []

        for block in paragraphs:
            block_labels = labels
            if label_forms:
                # Kemunculan pertama memakai bentuk panjang, berikutnya bentuk pendek.
                block_labels = {
                    index: (pendek if index in defined else panjang)
                    for index, (panjang, pendek) in label_forms.items()
                }
            # Spesifikasi "satu butir ditulis sebagai satu paragraf": blok yang
            # memuat butir bernomor/huruf dipecah satu paragraf per butir;
            # paragraf naratif biasa tetap utuh.
            groups = _numbered_groups(block)
            if groups is not None:
                texts = [
                    content.compose(content.Paragraph(group), block_labels)
                    for group in groups
                ]
            else:
                texts = [content.compose(block, block_labels)]
            for text in texts:
                if not text:
                    continue
                body = _insert_paragraph_before(paragraph, text, "Normal")
                # Tipografi akademik Indonesia: rata kiri-kanan, spasi 1,5,
                # dan baris pertama menjorok 1 cm tanpa jarak antarparagraf.
                body.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                body.paragraph_format.line_spacing = 1.5
                body.paragraph_format.first_line_indent = Cm(1.0)
                body.paragraph_format.space_after = Pt(0)
            if label_forms:
                for sentence in block.sentences:
                    if sentence.is_factual and sentence.citation:
                        defined.update(sentence.source_ids)
    _remove_paragraph(paragraph)


def _place_references(paragraph: DocxParagraph, references) -> None:
    """Tulis daftar pustaka; entri boleh string biasa atau daftar run (teks, italic).

    Run italic dipakai untuk nama jurnal + volume sesuai APA 7 (mis.
    *Healthcare, 9*(10), 1272.).
    """
    for entry in references:
        item = _insert_paragraph_before(paragraph, "", "Normal")
        item.paragraph_format.left_indent = _HANGING
        item.paragraph_format.first_line_indent = -_HANGING
        runs = entry if isinstance(entry, list) else [(str(entry), False)]
        previous = ""
        for text, italic in runs:
            text = str(text or "")
            if not text:
                continue
            # Nomor "(10)" menempel pada volume; segmen lain dipisah spasi.
            if previous and not text.startswith(("(", ",")):
                item.add_run(" ")
            run = item.add_run(placeholders.clean_text(text))
            if italic:
                run.italic = True
            previous = text
    _remove_paragraph(paragraph)


_TITLE_MINOR_WORDS = frozenset({
    "dan", "atau", "tetapi", "namun", "serta", "maupun", "yang", "di", "ke",
    "dari", "untuk", "pada", "dalam", "dengan", "sebagai", "terhadap", "bagi",
    "oleh", "karena", "agar", "antara", "selama", "hingga", "sampai", "tentang",
})


def _academic_title_case(title: str) -> str:
    """Kapitalisasi judul akademik formal (Title Case)."""
    raw = str(title or "").strip()
    if raw.startswith("{{") and raw.endswith("}}"):
        return raw
    cleaned = placeholders.clean_text(raw)
    words = cleaned.split()
    if not words:
        return cleaned
    letters = [c for c in cleaned if c.isalpha()]
    all_caps = bool(letters) and all(c.isupper() for c in letters)
    cased: list[str] = []
    for index, word in enumerate(words):
        core = word.strip('()[]:;,.!?"\'-')
        if not core:
            cased.append(word)
        elif len(core) >= 2 and core.isupper() and not all_caps:
            cased.append(word)  # akronim tetap kapital
        elif not all_caps and any(c.isupper() for c in core[1:]):
            cased.append(word)  # kata berkapital internal (LaporanX, iPhone) dipertahankan
        elif index > 0 and core.casefold() in _TITLE_MINOR_WORDS:
            cased.append(word.lower())
        else:
            lowered = word.lower()
            for pos, char in enumerate(lowered):
                if char.isalpha():
                    cased.append(lowered[:pos] + char.upper() + lowered[pos + 1:])
                    break
            else:
                cased.append(word)
    return " ".join(cased)


def render_report(
    template_path: Path,
    fields: dict[str, str],
    sections: list,
    references: list[str],
    logo_path: Path | None,
    output_path: Path,
    labels: dict[int, str] | None = None,
    update_fields: bool = False,
    label_forms: dict[int, tuple[str, str]] | None = None,
    add_page_numbers: bool = True,
) -> Path:
    doc = template.load_template(template_path)
    result = template.validate_document(doc)
    if not result.is_processable:
        detail: list[str] = []
        if result.missing:
            detail.append("penanda wajib tidak ada: " + ", ".join(result.missing))
        if result.not_standalone:
            detail.append(
                "penanda khusus harus berdiri sendiri: " + ", ".join(result.not_standalone)
            )
        raise template.TemplateError("Template tidak siap: " + "; ".join(detail))
    labels = labels or {}
    if update_fields:
        _enable_update_fields(doc)

    has_dosen = bool(str(fields.get("DOSEN") or "").strip() or str(fields.get("{{DOSEN}}") or "").strip())
    if not has_dosen:
        for p in list(placeholders.iter_paragraphs(doc)):
            p_text = p.text.strip().casefold()
            if p_text in ("dosen pengampu:", "dosen pembimbing:", "dosen pengampu", "dosen pembimbing"):
                _remove_paragraph(p)
            elif p_text == "{{dosen}}":
                _remove_paragraph(p)

    # Normalisasi dua arah: pemanggil boleh mengirim kunci "JUDUL" maupun
    # "{{JUDUL}}". Tanpa ini, kunci tanpa kurung kurawal tidak pernah
    # dicocokkan oleh replace_in_paragraph sehingga penanda tercetak mentah.
    mapping: dict[str, str] = {}
    for raw_key, raw_value in fields.items():
        key = str(raw_key).strip()
        if key.startswith("{{") and key.endswith("}}"):
            braced = key
            bare = key[2:-2].strip()
        else:
            braced = f"{{{{{key}}}}}"
            bare = key
        if not bare or braced in template.SPECIAL_PLACEHOLDERS:
            continue
        value = str(raw_value or "")
        if bare.upper() == "JUDUL":
            value = _academic_title_case(value)
        mapping[braced] = value
        mapping[bare] = value

    # Penanda opsional yang tidak dikirim pemanggil tetap dikosongkan agar
    # tidak pernah tercetak mentah di dokumen hasil (dua bentuk kunci).
    for name in template.OPTIONAL_PLACEHOLDERS:
        mapping.setdefault(name, "")
        mapping.setdefault(name[2:-2].strip(), "")

    # Penanda khusus diproses lebih dulu: nilai field bisa memuat teks yang
    # kebetulan sama dengan nama penanda khusus (mis. judul "{{ISI_LAPORAN}}").
    for paragraph in list(placeholders.iter_paragraphs(doc)):
        text = paragraph.text.strip()
        if text == "{{LOGO}}":
            _place_logo(paragraph, logo_path)
        elif text == "{{DAFTAR_ISI}}":
            _place_toc(paragraph)
        elif text == "{{ISI_LAPORAN}}":
            _place_body(doc, paragraph, sections, labels, label_forms)
        elif text == "{{DAFTAR_PUSTAKA}}":
            _place_references(paragraph, references)

    leftover = {name: "" for name in template.SPECIAL_PLACEHOLDERS}
    for paragraph in list(placeholders.iter_paragraphs(doc)):
        placeholders.replace_in_paragraph(paragraph, leftover)

    for paragraph in list(placeholders.iter_paragraphs(doc)):
        placeholders.replace_in_paragraph(paragraph, mapping)

    if add_page_numbers:
        _ensure_page_numbers(doc)

    final_path = _save_atomic(doc, output_path)
    return final_path


def _save_atomic(doc, output_path: Path) -> Path:
    """Simpan DOCX lewat berkas sementara; berkas lama tidak tersentuh bila gagal."""
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    descriptor, temporary = tempfile.mkstemp(
        prefix=f".{output_path.name}.", suffix=".docx", dir=output_path.parent
    )
    os.close(descriptor)
    target_path = output_path
    try:
        doc.save(temporary)
        try:
            os.replace(temporary, target_path)
        except PermissionError:
            timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
            target_path = output_path.parent / f"{output_path.stem}_{timestamp}.docx"
            os.replace(temporary, target_path)
        return target_path
    except BaseException:
        try:
            os.unlink(temporary)
        except OSError:
            pass
        raise
