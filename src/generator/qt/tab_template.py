"""Tab Template Qt: daftar DOCX dengan status validasi hitung-ulang (parity gui/tab_template.py).

Baru dibanding Tk: panel "Penanda terdeteksi" menampilkan hasil
`placeholders.collect_names` urut abjad untuk template terpilih. Dialog
metadata Tk diganti: `_import_template` menerima (src, nama, deskripsi)
lewat argumen agar bisa diuji tanpa dialog; bila nama None dipakai nama
berkas sumber.
"""

from __future__ import annotations

import shutil
from pathlib import Path

from PySide6.QtCore import Qt
from PySide6.QtWidgets import (
    QAbstractItemView,
    QFileDialog,
    QGroupBox,
    QHBoxLayout,
    QHeaderView,
    QLabel,
    QListWidget,
    QPushButton,
    QSplitter,
    QTableWidget,
    QTableWidgetItem,
    QVBoxLayout,
    QWidget,
)

from .. import paths
from ..database import Database, Template
from ..document.template import TemplateError, validate_template
from ..document import placeholders


class TemplateTab(QWidget):
    """Daftar template dengan status validasi hitung-ulang, Tambah, dan Hapus."""

    def __init__(self, db: Database):
        super().__init__()
        self._db = db
        self._items: list[Template] = []
        self._last_message = ""

        layout = QVBoxLayout(self)

        self._table = QTableWidget(0, 2)
        self._table.setHorizontalHeaderLabels(["Template", "Status"])
        self._table.horizontalHeader().setSectionResizeMode(QHeaderView.Stretch)
        self._table.setSelectionBehavior(QAbstractItemView.SelectRows)
        self._table.setSelectionMode(QAbstractItemView.SingleSelection)
        self._table.setEditTriggers(QAbstractItemView.NoEditTriggers)
        self._table.itemSelectionChanged.connect(self._on_select)
        self._detected = QListWidget()
        detected_box = QGroupBox("Penanda terdeteksi")
        detected_layout = QVBoxLayout(detected_box)
        detected_layout.addWidget(self._detected)

        splitter = QSplitter(Qt.Orientation.Horizontal)
        splitter.addWidget(self._table)
        splitter.addWidget(detected_box)
        splitter.setStretchFactor(0, 3)
        splitter.setStretchFactor(1, 2)
        layout.addWidget(splitter, stretch=1)

        buttons = QHBoxLayout()
        add = QPushButton("Tambah template (.docx)...")
        add.clicked.connect(self._add)
        delete = QPushButton("Hapus")
        delete.clicked.connect(self._delete)
        buttons.addWidget(add)
        buttons.addWidget(delete)
        buttons.addStretch(1)
        self._status = QLabel("")
        buttons.addWidget(self._status)
        layout.addLayout(buttons)

        self.refresh()

    def refresh(self) -> None:
        """Baca ulang daftar dan hitung status validasi dari berkasnya."""
        self._items = self._db.list_template()
        self._table.setRowCount(len(self._items))
        for row, item in enumerate(self._items):
            self._table.setItem(row, 0, QTableWidgetItem(item.nama))
            self._table.setItem(row, 1, QTableWidgetItem(self._status_text(item)))

    def _status_text(self, item: Template) -> str:
        path = Path(item.path)
        if not path.is_file():
            return "berkas hilang"
        try:
            result = validate_template(path)
        except TemplateError:
            return "rusak atau bukan .docx valid"
        problems = []
        if result.missing:
            problems.append(f"kurang {len(result.missing)} penanda wajib")
        if result.not_standalone:
            problems.append(f"{len(result.not_standalone)} penanda tidak berdiri sendiri")
        if result.unknown:
            problems.append(f"penanda tak dikenal: {', '.join(result.unknown)}")
        if not problems:
            return "lengkap"
        return "tidak lengkap (" + "; ".join(problems) + ")"

    def _selected(self) -> Template | None:
        row = self._table.currentRow()
        if row < 0 or row >= len(self._items):
            return None
        return self._items[row]

    def _on_select(self) -> None:
        item = self._selected()
        self._detected.clear()
        if item is None:
            return
        try:
            document = placeholders_document(Path(item.path))
        except Exception:
            return
        for name in sorted(placeholders.collect_names(document)):
            self._detected.addItem(name)

    def _detected_placeholder_texts(self) -> list[str]:
        return [self._detected.item(i).text() for i in range(self._detected.count())]

    def _add(self) -> None:
        chosen, _selected_filter = QFileDialog.getOpenFileName(
            self, "Pilih berkas template", "", "Dokumen Word (*.docx);;Semua berkas (*.*)"
        )
        if not chosen:
            return
        self._import_template(Path(chosen))

    def _import_template(self, src: Path, nama: str | None = None, deskripsi: str = "") -> bool:
        """Validasi, salin ke folder template, dan simpan metadata."""
        try:
            result = validate_template(src)
        except TemplateError as error:
            self._last_message = str(error)
            return False

        self._last_message = ""
        if result.missing or result.not_standalone or result.unknown:
            lines = []
            if result.missing:
                lines.append("Penanda wajib yang kurang: " + ", ".join(result.missing))
            if result.not_standalone:
                lines.append(
                    "Penanda khusus yang tidak berdiri sendiri: "
                    + ", ".join(result.not_standalone)
                )
            if result.unknown:
                lines.append("Penanda tak dikenal: " + ", ".join(result.unknown))
            lines.append("Template tetap disimpan dengan status tidak lengkap.")
            self._last_message = "\n".join(lines)

        if nama is None:
            nama = src.stem

        target = self._copy_into_templates_dir(src)
        if target is None:
            return False
        self._db.add_template(Template(0, nama.strip(), deskripsi.strip(), str(target)))
        self.refresh()
        self._status.setText(f"Template '{nama.strip()}' disimpan.")
        return True

    def _copy_into_templates_dir(self, src: Path) -> Path | None:
        paths.ensure_dirs()
        target = paths.TEMPLATES_DIR / src.name
        counter = 1
        while target.exists():
            target = paths.TEMPLATES_DIR / f"{src.stem}_{counter}{src.suffix}"
            counter += 1
        try:
            shutil.copy2(src, target)
        except OSError as error:
            self._last_message = f"Berkas template tidak bisa disalin: {error}"
            return None
        return target

    def _delete(self) -> None:
        item = self._selected()
        if item is None:
            self._status.setText("Pilih template yang akan dihapus.")
            return
        self._db.delete_template(item.id)
        try:
            Path(item.path).unlink(missing_ok=True)
        except OSError:
            pass
        self.refresh()
        self._status.setText(f"Template '{item.nama}' dihapus.")


def placeholders_document(path: Path):
    """Buka DOCX untuk dibaca penandanya (kembali Document python-docx)."""
    from docx import Document

    return Document(path)
