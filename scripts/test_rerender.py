import sys
sys.path.insert(0, 'src')
from pathlib import Path
from generator.drafts import load_draft
from generator.document import renderer, template
from generator import report, format_spec, reviewed
import docx

draft = load_draft("aktif")
if not draft:
    print("Draft aktif not found!")
    sys.exit(1)

doc = reviewed.ReviewedDocument.from_payload(draft.sections, draft.sources)
print("Loaded draft aktif:")
print(" - Title in draft fields:", repr(draft.fields.get("JUDUL")))
print(" - Dosen in draft fields:", repr(draft.fields.get("DOSEN")))
print(" - Variant:", draft.fields.get("VARIANT"))

v = format_spec.get_variant(draft.fields.get("VARIANT", "D"))
norm_sections = format_spec.normalize_sections(v, doc.sections)

print("\nNormalized sections in Bab 3:")
for s in norm_sections:
    if "3." in s.title or "PEMBAHASAN" in s.title:
        print("  -", s.title)

cited = doc.cited_ids
ref_runs = report.reference_runs_for(doc.sources, cited)
labels = report.labels_for(doc.sources, cited)

out_path = Path("output/dampak_inovasi_teknologi_pada_pertanian_REVISED.docx")
tpl_path = Path("data/templates/Template_Standar_Akademik.docx")

final_path = renderer.render_report(
    template_path=tpl_path,
    fields=draft.fields,
    sections=norm_sections,
    references=ref_runs,
    logo_path=Path(draft.fields.get("LOGO_PATH", "")) if draft.fields.get("LOGO_PATH") else None,
    output_path=out_path,
    labels=labels,
    update_fields=True,
)

print(f"\nRendered revised docx to: {final_path}")

# Now inspect the rendered docx
out_doc = docx.Document(final_path)
print("\n=== COVER PARAGRAPHS ===")
for i in range(min(12, len(out_doc.paragraphs))):
    print(f"[{i}] {out_doc.paragraphs[i].text}")

print("\n=== DAFTAR PUSTAKA PARAGRAPHS ===")
dp_start = False
for i, p in enumerate(out_doc.paragraphs):
    if "DAFTAR PUSTAKA" in p.text:
        dp_start = True
    elif dp_start and p.text.strip():
        print(f"[{i}] {p.text}")

print("\n=== BAB 1 PARAGRAPH 1 (CITATION SMOOTHING CHECK) ===")
for p in out_doc.paragraphs:
    if "Sistem pertanian global saat ini" in p.text:
        print("Text:\n", p.text)
