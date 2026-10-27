"""Pemeriksa E2E run-nyata — Penjaga Regresi (Fase 0b).

Menjalankan pipeline penuh lewat ApiBridge dengan layanan AI sungguhan
(add source tersimpan + start_generation varian "b" + ekspor DOCX), lalu
mencetak PASS/FAIL per invariant beserta metrik (kata, paragraf/subbab, % hijau).

Pemakaian:  .venv/Scripts/python.exe scripts/e2e_check.py
Data pengguna tidak disentuh: basis data, output, dan draf diarahkan ke
folder sementara di bawah scratch/.
"""
from __future__ import annotations

import json
import re
import sys
import time
from datetime import datetime
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))

import docx  # noqa: E402

from generator import database, format_spec, paths  # noqa: E402
from generator.ai.content import _token_similarity  # noqa: E402
from generator.web import api_bridge  # noqa: E402

TOPIC = "Pengaruh Media Sosial terhadap Prestasi Akademik Mahasiswa"
RUN_DIR = ROOT / "scratch" / "e2e_runs" / datetime.now().strftime("%Y%m%d-%H%M%S")
SOURCES_FILE = ROOT / "scratch" / "student_run" / "final_sources.json"
NUMBERED = re.compile(r"^\s*(?:\d{1,2}|[a-z])[\.\)]\s+\S")
MARKER = re.compile(r"\[(?:TABEL|GAMBAR|FOTO)\s*:", re.IGNORECASE)
LIST_TITLES = ("rumusan masalah", "tujuan", "kesimpulan", "saran")

# --- Guard S3 (Fase 2) ------------------------------------------------------
# Kata fungsi & penanda klausa diimpor dari format_spec (jangan disalin); judul
# baku non-dinamis varian "b" dikecualikan dari aturan minimal 3 kata karena
# judul seperti "Hasil"/"Saran" memang sah 1 kata.
_FUNCTION_WORDS = format_spec._FUNCTION_WORDS
_CLAUSE_MARKERS = format_spec._CLAUSE_MARKERS
_MAX_TITLE_WORDS = 10
_TRAILING_PUNCT = ".,;:!?\"'-–—"
_CANONICAL_TITLE_KEYS = {
    format_spec._key(entry.title)
    for entry in format_spec._entries(format_spec.get_variant("b"))
    if entry.level >= 2 and not entry.node.dynamic
}


def _title_without_number(title: str) -> str:
    """Judul tanpa nomor subbab: hitungan kata & normalisasi memakai bentuk ini."""
    return re.sub(r"^[\d.]+\s*", "", str(title or "")).strip()


def _chapter_token(title: str) -> str:
    """Nomor bab dari token pertama judul ("2.2.1" -> "2"); "" bila bukan nomor."""
    first = str(title or "").split()[0] if str(title or "").split() else ""
    return first.split(".")[0] if re.fullmatch(r"[\d.]+", first) else ""


def _s3_problems(sections: list[dict]) -> list[str]:
    """Masalah judul (S3): kata berlebih, kata fungsi di ujung, kembar selevel.

    Cakupan khusus heading level >= 2. Butir (a)-(e) sama dengan guard di
    tests/test_e2e_invariants.py supaya run-nyata dan tes deterministik setara.
    """
    problems: list[str] = []
    seen: dict[tuple[str, int], dict[str, str]] = {}
    for section in sections:
        level = int(section.get("level", 1))
        if level < 2:
            continue
        title = str(section.get("title", "") or "").strip()
        stripped = _title_without_number(title)
        words = stripped.split()
        key = format_spec._key(stripped)
        canonical = key in _CANONICAL_TITLE_KEYS
        if len(words) > _MAX_TITLE_WORDS:
            problems.append(f"judul > 10 kata: {title[:50]}")
        if words and words[-1].strip(_TRAILING_PUNCT).casefold() in _FUNCTION_WORDS:
            problems.append(f"judul berakhir kata fungsi: {title[:50]}")
        if not canonical and any(
            word.strip(_TRAILING_PUNCT).casefold() in _CLAUSE_MARKERS for word in words
        ):
            problems.append(f"judul memuat penanda klausa: {title[:50]}")
        if not canonical and len(words) < 3:
            problems.append(f"judul < 3 kata non-kanonik: {title[:50]}")
        if re.match(
            r"^(berbagai\s+(penelitian|studi)|banyak\s+(penelitian|studi)|bagaimana|mengapa|"
            r"penelitian\s+ini|hasil\s+penelitian)",
            stripped,
            re.I,
        ):
            problems.append(f"judul tampak kalimat: {title[:50]}")
        group = seen.setdefault((_chapter_token(title), level), {})
        if key and key in group:
            problems.append(f"judul kembar selevel '{title[:40]}' dengan '{group[key][:40]}'")
        group.setdefault(key, title)
    return problems



def _sources() -> list[dict]:
    if SOURCES_FILE.is_file():
        return json.loads(SOURCES_FILE.read_text(encoding="utf-8"))
    raise SystemExit(
        f"Sumber uji tidak ditemukan: {SOURCES_FILE}. Jalankan dulu "
        "scratch/student_run/step4_generate.py atau sediakan sumber manual."
    )


def _generate(bridge, sources: list[dict]) -> None:
    started = bridge.start_generation({
        "topic": TOPIC, "outline": [], "sources": sources, "variant": "b",
    })
    if not started.get("started") and "error" in started:
        raise SystemExit(f"start_generation gagal: {started}")
    while bridge._gen_thread and bridge._gen_thread.is_alive():
        time.sleep(2)


def _analyze_docx(path: Path) -> dict:
    document = docx.Document(str(path))
    headings: list[dict] = []
    paragraphs: list[str] = []
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue
        if paragraph.style.name.startswith("Heading"):
            headings.append({"title": text, "style": paragraph.style.name, "body": []})
        elif headings:
            headings[-1]["body"].append(text)
        else:
            paragraphs.append(text)
    return {
        "headings": headings,
        "cover": paragraphs,
        "all_text": [p.text.strip() for p in document.paragraphs if p.text.strip()],
        "document": document,
    }


def _section_texts(sections: list[dict]) -> list[str]:
    return [
        sentence.get("text", "").strip()
        for section in sections
        for paragraph in section.get("paragraphs", [])
        for sentence in paragraph.get("sentences", [])
        if sentence.get("text", "").strip()
    ]


def _check(sections: list[dict], statuses: list[dict], docx_info: dict, references: list[str]) -> dict:
    results: dict[str, bool] = {}
    notes: dict[str, str] = {}
    titles = [section.get("title", "") for section in sections]

    def fail(key: str, message: str) -> None:
        results[key] = False
        notes.setdefault(key, message)

    def ok(key: str) -> None:
        results.setdefault(key, True)

    for key in ("S1", "S2", "S3", "S4", "S5", "S6", "S7", "K1", "K2", "K3", "K4", "V1", "V2", "V3", "E1"):
        ok(key)

    # S1: BAB 2 punya subbab dan heading bab tanpa isi
    bab2 = [s for s in sections if s.get("title", "").startswith("BAB 2.")]
    subs2 = [s for s in sections if int(s.get("level", 1)) >= 2 and s.get("title", "").startswith("2.")]
    if not bab2 or bab2[0].get("paragraphs"):
        fail("S1", "heading BAB 2 masih memuat paragraf")
    if not any(s["title"].startswith("2.1") for s in subs2) or not any(
        s["title"].startswith("2.2") for s in subs2
    ):
        fail("S1", f"subbab 2.1/2.2 hilang: {[s['title'] for s in subs2][:6]}")
    if not any(s["title"].startswith("2.1") for s in subs2):
        ok("S2")  # nomor tidak bisa dicek tanpa subbab

    # S2: penomoran rapat
    level2_by_chapter: dict[str, list[int]] = {}
    level3_by_parent: dict[str, list[int]] = {}
    for section in sections:
        number = section.get("title", "").split(" ")[0]
        if not re.match(r"^[\d.]+$", number):
            continue
        parts = number.rstrip(".").split(".")
        if len(parts) == 2 and int(section.get("level", 1)) == 2:
            level2_by_chapter.setdefault(parts[0], []).append(int(parts[1]))
        elif len(parts) == 3:
            level3_by_parent.setdefault(".".join(parts[:2]), []).append(int(parts[2]))
    for chapter, numbers in level2_by_chapter.items():
        if sorted(numbers) != list(range(1, len(numbers) + 1)):
            fail("S2", f"nomor level 2 bab {chapter} tidak rapat: {sorted(numbers)}")
    for parent, numbers in level3_by_parent.items():
        if sorted(numbers) != list(range(1, len(numbers) + 1)):
            fail("S2", f"nomor level 3 {parent} tidak rapat: {sorted(numbers)}")
    # S2 tambahan: urutan fisik sekuensial bab wajib monoton naik (tidak boleh ada subbab Bab 2 sebelum Bab 1)
    curr_chap = 0
    first_level1_seen = False
    for section in sections:
        lvl = int(section.get("level", 1))
        if lvl == 1:
            first_level1_seen = True
        elif not first_level1_seen:
            fail("S2", f"subbab terdeteksi sebelum BAB 1: {section.get('title', '')[:35]}")
        nt = format_spec._number_tuple(section.get("title", ""))
        chap_num = nt[0]
        if chap_num not in (0, 998, 999):
            if chap_num < curr_chap:
                fail("S2", f"inversi urutan bab: {section.get('title', '')[:35]} setelah Bab {curr_chap}")
            curr_chap = max(curr_chap, chap_num)

    # S3: judul frasa nomina pendek, tanpa kata fungsi menggantung, unik selevel
    for problem in _s3_problems(sections):
        fail("S3", problem)

    # S4: tanpa judul template/pengisi
    if any("{" in title or "}" in title for title in titles):
        fail("S4", "judul template {...} tersisa")
    if any(re.search(r"\b(pengisi|tambahan|topik \d)\b", title, re.I) for title in titles):
        fail("S4", "judul pengisi terdeteksi")
    filler = [title for title in titles if title.split()[0].count(".") and title.split()[0].split(".")[1] in ("3", "4") and any(
        title.startswith(prefix) for prefix in ("2.3", "2.4", "2.5", "2.6", "2.7", "2.8", "2.9", "2.10", "2.11", "2.12", "2.13")
    )]
    if filler:
        fail("S4", f"subbab pengisi gaya 2.3+: {filler[:3]}")

    # S5: BAB 4 hanya butir
    bab4 = [s for s in sections if s.get("title", "").startswith("BAB 4.")]
    if bab4 and bab4[0].get("paragraphs"):
        fail("S5", "paragraf menempel di level BAB 4")
    if not any(s["title"].startswith("4.1") for s in sections) or not any(
        s["title"].startswith("4.2") for s in sections
    ):
        fail("S5", "subbab 4.1/4.2 hilang")

    # S6: satu butir satu paragraf (dari DOCX)
    numbered_paragraphs = [text for text in docx_info["all_text"] if NUMBERED.match(text)]
    if any(len(NUMBERED.findall(text)) > 1 for text in docx_info["all_text"]):
        fail("S6", "beberapa butir bernomor dalam satu paragraf")
    if not numbered_paragraphs:
        fail("S6", "tidak ada paragraf butir bernomor")

    # S7/K5: kedalaman
    latar = next((s for s in sections if "Latar Belakang" in s.get("title", "")), None)
    if latar is None or len(latar.get("paragraphs", [])) < 2:
        fail("S7", "1.1 Latar Belakang < 2 paragraf")
    for section in sections:
        title = section.get("title", "")
        if int(section.get("level", 1)) < 2 or any(k in title.casefold() for k in LIST_TITLES):
            continue
        paragraphs = section.get("paragraphs", [])
        if not paragraphs:
            continue
        if len(paragraphs) > 6:
            fail("S7", f"{title[:40]} > 6 paragraf")
        if title.startswith(("2.", "3.")) and len(paragraphs) < 2:
            fail("S7", f"{title[:40]} < 2 paragraf")
        if int(section.get("level", 1)) == 2 and title.startswith("3.") and len(paragraphs) < 3:
            fail("K5", f"{title[:40]} < 3 paragraf")
        for paragraph in paragraphs:
            count = len(paragraph.get("sentences", []))
            if count > 5:
                fail("S7", f"{title[:40]} paragraf {count} kalimat (>5)")

    texts = _section_texts(sections)
    # K1: tanpa penanda tabel/deskriptor
    if any(MARKER.search(text) for text in texts):
        fail("K1", "penanda [TABEL/GAMBAR/FOTO] tersisa")
    if any(re.search(r"(penanda\s+tabel|disajikan\s+dalam\s+bentuk\s+tabel)", text, re.I) for text in texts):
        fail("K1", "kalimat deskriptor tabel tersisa")

    # K2: kalimat yatim huruf kecil (kecuali butir daftar bertanda huruf seperti a., b.)
    for text in texts:
        t = text.strip()
        if not t:
            continue
        # Buang HANYA nomor + tanda baca + spasi; jangan gunakan NUMBERED.sub
        # karena \S di ujungnya ikut memakan huruf pertama kata.
        core = re.sub(r"^\s*(?:\d{1,2}|[a-z])[\.\)]\s+", "", t).strip()
        if core and core[:1].islower():
            fail("K2", f"butir bernomor diawali huruf kecil: {t[:40]}")
        elif not core and t[:1].islower():
            fail("K2", f"kalimat diawali huruf kecil: {t[:40]}")
        elif not NUMBERED.match(t) and t[:1].islower():
            fail("K2", f"kalimat diawali huruf kecil: {t[:40]}")

    # K3: nomor butir unik
    for section in sections:
        if not any(k in section.get("title", "").casefold() for k in LIST_TITLES):
            continue
        numbers = [
            NUMBERED.match(sentence["text"].strip()).group(0).strip()
            for paragraph in section.get("paragraphs", [])
            for sentence in paragraph.get("sentences", [])
            if NUMBERED.match(sentence.get("text", "").strip())
        ]
        if len(numbers) != len(set(numbers)):
            fail("K3", f"nomor butir ganda di {section.get('title','')[:30]}")

    # K4: duplikat identik & parafrase
    seen: dict[str, str] = {}
    for section in sections:
        for text in [
            sentence.get("text", "").strip()
            for paragraph in section.get("paragraphs", [])
            for sentence in paragraph.get("sentences", [])
        ]:
            key = " ".join(text.lower().split())
            if key in seen and seen[key] != section.get("title"):
                fail("K4", "duplikat identik lintas bab")
            seen.setdefault(key, section.get("title"))
    for section in sections:
        if int(section.get("level", 1)) < 2:
            continue
        sentences = [
            sentence.get("text", "")
            for paragraph in section.get("paragraphs", [])
            for sentence in paragraph.get("sentences", [])
        ]
        for index in range(len(sentences)):
            for other in range(index + 1, len(sentences)):
                if _token_similarity(sentences[index], sentences[other]) >= 0.55:
                    fail("K4", f"parafrase ganda di {section.get('title','')[:30]}")

    # V1: sitasi & pustaka
    cited = {
        source_id
        for section in sections
        for paragraph in section.get("paragraphs", [])
        for sentence in paragraph.get("sentences", [])
        if sentence.get("citation")
        for source_id in sentence.get("source_ids", [])
    }
    if not cited:
        fail("V1", "tidak ada sitasi dalam teks")
    if len(references) != len(cited):
        fail("V1", f"jumlah pustaka {len(references)} != sumber dikutip {len(cited)}")

    # V2: volume(nomor), halaman, italic
    joined_refs = " ".join(references)
    if not re.search(r"\d+\(\d+\),", joined_refs):
        fail("V2", "referensi tidak memuat volume(nomor)")
    italic_runs = [
        run.text
        for paragraph in docx_info["document"].paragraphs
        for run in paragraph.runs
        if run.italic
    ]
    if not italic_runs:
        fail("V2", "tidak ada run italic pada referensi")

    # V3: proporsi hijau & terjemahan
    total = len(statuses)
    green = sum(1 for status in statuses if status.get("status") == "terverifikasi")
    ratio = green / total if total else 0.0
    if ratio < 0.8:
        fail("V3", f"hijau hanya {green}/{total} ({ratio * 100:.0f}%)")

    # E1: ekspor bersih
    if any("{{" in text for text in docx_info["all_text"]):
        fail("E1", "penanda {{...}} tersisa")
    if not any("TOC" in paragraph._p.xml for paragraph in docx_info["document"].paragraphs):
        fail("E1", "field TOC tidak ada")
    if "PAGE" not in docx_info["document"].sections[0].footer._element.xml:
        fail("E1", "nomor halaman tidak ada")
    cover = "\n".join(docx_info["cover"][:20])
    if "PROGRAM STUDI" not in cover.upper() or "FAKULTAS" not in cover.upper():
        fail("E1", "sampul prodi/fakultas tidak terisi")

    words = sum(len(text.split()) for text in texts)
    return {"results": results, "notes": notes, "words": words, "green_ratio": ratio,
            "total_sentences": total, "green": green}


def main() -> int:
    RUN_DIR.mkdir(parents=True, exist_ok=True)
    real_config = paths.CONFIG_PATH  # konfigurasi layanan AI milik pengguna (baca saja)
    paths.set_data_dir(RUN_DIR / "data")
    paths.OUTPUT_DIR = RUN_DIR / "output"
    paths.OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

    db = database.Database(RUN_DIR / "check.db")
    db.init_schema()
    bridge = api_bridge.ApiBridge(db, config_path=real_config)
    bridge.save_profile({
        "nama": "Raka Pratama", "nim": "241510101001",
        "prodi": "Agroteknologi", "fakultas": "Pertanian",
        "institusi": "Universitas Jember",
    })
    if not bridge.list_templates():
        bridge.create_default_template({"font_name": "Times New Roman", "margin_preset": "4433"})
    template_id = bridge.list_templates()[0]["id"]

    sources = _sources()
    print(f"[e2e] sumber: {len(sources)} | model: {bridge._cfg.ai_model}")
    _generate(bridge, sources)

    draft = bridge.load_active_draft()
    if not draft.get("has_draft"):
        print("FAIL: draf tidak terbentuk")
        return 1
    sections = draft["sections"]
    statuses = bridge.verify_document(sections, draft["sources"])
    fields = dict(draft.get("fields", {}))
    fields.update({
        "NAMA": "Raka Pratama", "NIM": "241510101001",
        "PRODI": "Agroteknologi", "FAKULTAS": "Pertanian",
        "INSTITUSI": "Universitas Jember",
        "MATA_KULIAH": "Metodologi Penelitian",
        "DOSEN": "Dr. Ir. Soekarno Hadi, M.P.",
    })
    state = {
        "fields": fields, "sources": draft["sources"], "sections": sections,
        "done": draft.get("done", []), "template_id": template_id, "open_word": False,
    }
    preflight = bridge.run_preflight(state)
    export = bridge.export_report(state)
    if not export.get("success"):
        print("FAIL ekspor:", export.get("error"))
        return 1
    references: list[str] = []
    docx_info = _analyze_docx(Path(export["output_path"]))
    # Fase 2: cetak semua heading DOCX (satu baris per heading) agar kualitas
    # judul bisa diperiksa sekali lihat tanpa membuka berkas Word.
    for heading in docx_info["headings"]:
        print(f"[e2e] heading {heading['style']}: {heading['title']}")
    bibliography_index = next(
        (index for index, heading in enumerate(docx_info["headings"])
         if heading["title"].upper() == "DAFTAR PUSTAKA"),
        None,
    )
    if bibliography_index is not None:
        references = docx_info["headings"][bibliography_index]["body"]

    report = _check(sections, statuses, docx_info, references)
    print(f"[e2e] kata={report['words']} kalimat={report['total_sentences']} "
          f"hijau={report['green']}/{report['total_sentences']} "
          f"({report['green_ratio'] * 100:.0f}%)")
    print(f"[e2e] preflight issues: {[issue['kind'] for issue in preflight['issues']]}")
    for key, passed in report["results"].items():
        mark = "PASS" if passed else "FAIL"
        detail = "" if passed else f" — {report['notes'].get(key, '')}"
        print(f"  [{mark}] {key}{detail}")
    print(f"[e2e] DOCX: {export['output_path']}")
    return 0 if all(report["results"].values()) else 1


if __name__ == "__main__":
    raise SystemExit(main())